import re
import os
import sys
import json
import pptx
from datetime import datetime

# Reconfigure stdout to use UTF-8
sys.stdout.reconfigure(encoding='utf-8')

def load_dotenv():
    """Load environment variables from .env file if it exists."""
    if os.path.exists(".env"):
        with open(".env", "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if "=" in line and not line.startswith("#"):
                    k, v = line.split("=", 1)
                    os.environ[k.strip()] = v.strip()

load_dotenv()

def parse_pptx_items(pptx_path):
    """Parse slide items from PPTX with the sub-item heuristic and handle duplicate numbers."""
    prs = pptx.Presentation(pptx_path)
    parsed_items = []
    
    for slide_idx, slide in enumerate(prs.slides, 1):
        slide_text_blocks = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if "☆" in text:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    p_text = paragraph.text.strip()
                    if p_text:
                        slide_text_blocks.append(p_text)
                        
        current_num = None
        current_lines = []
        
        def flush_item():
            nonlocal current_num, current_lines
            if current_num is not None and current_lines:
                text = " ".join(current_lines).strip()
                text = re.sub(r"\s+", " ", text)
                parsed_items.append({
                    "slide": slide_idx,
                    "num": current_num,
                    "text": text
                })
            current_lines = []
            current_num = None

        for block in slide_text_blocks:
            match = re.match(r"^(\d+)\s*\.\s*(.*)", block, re.DOTALL)
            if match:
                num = int(match.group(1))
                text = match.group(2)
                # Heuristic: if current_num exists and the new number is much smaller, it's a subitem
                if current_num is not None and (current_num - num > 20):
                    current_lines.append(block)
                else:
                    flush_item()
                    current_num = num
                    current_lines.append(text)
            else:
                if current_num is not None:
                    current_lines.append(block)
                else:
                    lines = block.split('\n')
                    for line in lines:
                        line = line.strip()
                        m = re.match(r"^(\d+)\s*\.\s*(.*)", line)
                        if m:
                            num = int(m.group(1))
                            t = m.group(2)
                            if current_num is not None and (current_num - num > 20):
                                current_lines.append(line)
                            else:
                                flush_item()
                                current_num = num
                                current_lines.append(t)
                        elif current_num is not None:
                            current_lines.append(line)
        flush_item()

    # Suffix duplicate numbers
    num_counts = {}
    for item in parsed_items:
        num = item['num']
        num_counts[num] = num_counts.get(num, 0) + 1
        
    seen_counts = {}
    for item in parsed_items:
        num = item['num']
        if num_counts[num] > 1:
            seen_counts[num] = seen_counts.get(num, 0) + 1
            item['key'] = f"{num}-{seen_counts[num]}"
        else:
            item['key'] = str(num)
            
    return parsed_items

def parse_existing_md(md_path):
    """Parse Part 1 (Principles 1-20), chapters with golden quotes, and other metadata from markdown."""
    if not os.path.exists(md_path):
        return "", [], []
        
    with open(md_path, "r", encoding="utf-8-sig") as f:
        content = f.read()
        
    parts = content.split("## 第二部分：簡報對照與逐條潤稿 (簡報版)")
    part1_content = parts[0]
    
    # Parse chapters and their golden quotes
    chapters = []
    chap_matches = list(re.finditer(r"## (第[一二三四五六]篇[^\n]+)", part1_content))
    for i, m in enumerate(chap_matches):
        chap_title = m.group(1).strip()
        start_pos = m.end()
        end_pos = chap_matches[i+1].start() if i+1 < len(chap_matches) else len(part1_content)
        chap_body = part1_content[start_pos:end_pos]
        quotes = re.findall(r">\s*-\s*\*\*☆(.*?)☆\*\*", chap_body)
        chapters.append({
            "id": str(i+1),
            "title": chap_title,
            "quotes": [f"☆{q.strip()}☆" for q in quotes]
        })

    # Parse principles
    principle_pattern = re.compile(r"#### 原則 (\d+)：(.*?)\n(.*?)(?=\n#### 原則|\n##|\n---|\Z)", re.DOTALL)
    principles = []
    
    for match in principle_pattern.finditer(part1_content):
        num = int(match.group(1))
        title = match.group(2).strip()
        body = match.group(3).strip()
        
        mandarin = ""
        english = ""
        taiwanese = ""
        ref = ""
        
        m_match = re.search(r"\*\s+\*\*國語 \(Mandarin\)\*\*：(.*?)(?=\n\*|\n\[對照|\Z)", body)
        e_match = re.search(r"\*\s+\*\*English\*\*：(.*?)(?=\n\*|\n\[對照|\Z)", body)
        t_match = re.search(r"\*\s+\*\*台語 \(Taiwanese\)\*\*：(.*?)(?=\n\*|\n\[對照|\Z)", body)
        ref_match = re.search(r"`\[對照原項目：(.*?)\]`", body)
        
        if m_match: mandarin = m_match.group(1).strip()
        if e_match: english = e_match.group(1).strip()
        if t_match: taiwanese = t_match.group(1).strip()
        if ref_match: ref = ref_match.group(1).strip()
        
        # Clean ref brackets if any
        ref = ref.replace("[對照原項目：", "").replace("]", "").strip()
        
        principles.append({
            "num": num,
            "title": title,
            "mandarin": mandarin,
            "english": english,
            "taiwanese": taiwanese,
            "ref": ref
        })
        
    return part1_content, principles, chapters

def translate_via_gemini(original_text):
    """Translate original English text into Mandarin, English, and Taiwanese using Gemini API."""
    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key:
        return None
        
    try:
        from google import genai
        client = genai.Client(api_key=api_key)
        
        prompt = f"""
You are an expert translator and copywriter specializing in Taiwanese Hokkien (台語漢字) and professional business context.
Please translate and polish the following workplace principle:
"{original_text}"

You must return a JSON object with exactly the following keys:
- "english": Polished, professional English version.
- "mandarin": Polished, natural Traditional Chinese (Mandarin/國語) version suitable for business context.
- "taiwanese": Idiomatic Taiwanese Hokkien (台語漢字) version written in standard Traditional Chinese characters representing spoken Taiwanese, capturing the business wisdom naturally (e.g. use standard terms like "鬥相共", "省氣力", "做代誌", etc.).

Return ONLY the raw JSON string without markdown blocks or explanation.
"""
        model_name = os.environ.get("GEMINI_MODEL", "gemini-3.8-flash")
        response = client.models.generate_content(
            model=model_name,
            contents=prompt,
        )
        
        # Parse JSON from response
        res_text = response.text.strip()
        # strip markdown formatting if the model returned it
        if res_text.startswith("```"):
            res_text = re.sub(r"^```(?:json)?\n", "", res_text)
            res_text = re.sub(r"\n```$", "", res_text)
            
        data = json.loads(res_text.strip())
        return data
    except Exception as e:
        print(f"Error calling Gemini API: {e}")
        return None

def generate_odt_handbook(principles, final_items, odt_path, chapters=None):
    """Generate a well-formatted ODT document for the handbook ready for printing/book publishing."""
    from odf.opendocument import OpenDocumentText
    from odf.style import Style, TextProperties, ParagraphProperties, PageLayout, PageLayoutProperties, MasterPage
    from odf.text import H, P, Span

    print("\nStep 7: Generating ODT handbook for publishing...")
    doc = OpenDocumentText()
    
    # ------------------ DEFINE PAGE LAYOUT (A5) ------------------
    page_layout = PageLayout(name="A5Layout")
    page_layout.addElement(PageLayoutProperties(
        pagewidth="14.8cm",
        pageheight="21.0cm",
        marginleft="1.5cm",
        marginright="1.5cm",
        margintop="1.5cm",
        marginbottom="1.5cm"
    ))
    doc.automaticstyles.addElement(page_layout)
    
    master_page = MasterPage(name="Standard", pagelayoutname=page_layout)
    doc.masterstyles.addElement(master_page)
    
    # ------------------ DEFINE STYLES ------------------
    # Title style
    title_style = Style(name="BookTitle", family="paragraph")
    title_style.addElement(TextProperties(fontsize="26pt", fontweight="bold", fontfamily="Microsoft JhengHei", color="#1e293b"))
    title_style.addElement(ParagraphProperties(textalign="center", margintop="2cm", marginbottom="0.5cm"))
    doc.styles.addElement(title_style)
    
    # Subtitle style
    subtitle_style = Style(name="BookSubtitle", family="paragraph")
    subtitle_style.addElement(TextProperties(fontsize="14pt", fontfamily="Microsoft JhengHei", color="#4f46e5"))
    subtitle_style.addElement(ParagraphProperties(textalign="center", margintop="0.2cm", marginbottom="3cm"))
    doc.styles.addElement(subtitle_style)
    
    # Meta style (author, date etc.)
    meta_style = Style(name="BookMeta", family="paragraph")
    meta_style.addElement(TextProperties(fontsize="10pt", fontfamily="Microsoft JhengHei", color="#64748b"))
    meta_style.addElement(ParagraphProperties(textalign="center", margintop="5cm", marginbottom="1cm"))
    doc.styles.addElement(meta_style)
    
    # Chapter Heading style (H1 with Page Break, Centered for Divider Page)
    chapter_style = Style(name="BookChapter", family="paragraph")
    chapter_style.addElement(TextProperties(fontsize="20pt", fontweight="bold", fontfamily="Microsoft JhengHei", color="#4f46e5"))
    chapter_style.addElement(ParagraphProperties(breakbefore="page", textalign="center", margintop="4.5cm", marginbottom="1cm"))
    doc.styles.addElement(chapter_style)
    
    # Principle Heading style (H2 with Page Break)
    principle_title_style = Style(name="BookPrincipleTitle", family="paragraph")
    principle_title_style.addElement(TextProperties(fontsize="14pt", fontweight="bold", fontfamily="Microsoft JhengHei", color="#1e293b"))
    principle_title_style.addElement(ParagraphProperties(breakbefore="page", margintop="0.8cm", marginbottom="0.4cm", keepwithnext="true"))
    doc.styles.addElement(principle_title_style)
    
    # Principle Translation Body style
    principle_body_style = Style(name="BookPrincipleBody", family="paragraph")
    principle_body_style.addElement(TextProperties(fontsize="10.5pt", fontfamily="Microsoft JhengHei", color="#334155"))
    principle_body_style.addElement(ParagraphProperties(margintop="0.1cm", marginbottom="0.1cm", lineheight="150%"))
    doc.styles.addElement(principle_body_style)
    
    # Bold inline label style
    bold_label_style = Style(name="BoldLabel", family="text")
    bold_label_style.addElement(TextProperties(fontweight="bold", color="#1e293b"))
    doc.styles.addElement(bold_label_style)
    
    # Slide Item Heading style (H3 without page break)
    slide_title_style = Style(name="BookSlideTitle", family="paragraph")
    slide_title_style.addElement(TextProperties(fontsize="11pt", fontweight="bold", fontfamily="Microsoft JhengHei", color="#0f766e")) # Teal
    slide_title_style.addElement(ParagraphProperties(margintop="0.6cm", marginbottom="0.2cm", keepwithnext="true"))
    doc.styles.addElement(slide_title_style)
    
    # Slide Item Heading style with page break (H3 with page break)
    slide_title_style_with_break = Style(name="BookSlideTitleWithBreak", family="paragraph")
    slide_title_style_with_break.addElement(TextProperties(fontsize="11pt", fontweight="bold", fontfamily="Microsoft JhengHei", color="#0f766e")) # Teal
    slide_title_style_with_break.addElement(ParagraphProperties(breakbefore="page", margintop="0.6cm", marginbottom="0.2cm", keepwithnext="true"))
    doc.styles.addElement(slide_title_style_with_break)
    
    # Slide Item Original style (Monospace)
    slide_orig_style = Style(name="BookSlideOrig", family="paragraph")
    slide_orig_style.addElement(TextProperties(fontsize="9.5pt", fontfamily="Courier New", color="#475569", fontstyle="italic"))
    slide_orig_style.addElement(ParagraphProperties(margintop="0.1cm", marginbottom="0.1cm", marginleft="0.5cm"))
    doc.styles.addElement(slide_orig_style)
    
    # Slide Item Translation style
    slide_trans_style = Style(name="BookSlideTrans", family="paragraph")
    slide_trans_style.addElement(TextProperties(fontsize="9.5pt", fontfamily="Microsoft JhengHei", color="#334155"))
    slide_trans_style.addElement(ParagraphProperties(margintop="0.1cm", marginbottom="0.1cm", marginleft="0.5cm", lineheight="140%"))
    doc.styles.addElement(slide_trans_style)
    
    # ------------------ BUILD COVER PAGE ------------------
    p = P(stylename=title_style)
    p.addText("工作的管見 (My Two Cents)")
    doc.text.addElement(p)
    
    p = P(stylename=subtitle_style)
    p.addText("From the Floor: My Two Cents on Practical Work Principles\n職場生存與成長的避坑指南 (三語對照手冊)")
    doc.text.addElement(p)
    
    p = P(stylename=meta_style)
    p.addText("精裝出書版面格式\n© 2026 工作的管見")
    doc.text.addElement(p)
    
    # ------------------ GENERATE CHAPTERS & PRINCIPLES ------------------
    def get_chapter_name(num):
        if num <= 4: return '第一篇：效能與工具（生產力管理）'
        if num <= 8: return '第二篇：思考與決策（思維模型與問題解決）'
        if num <= 11: return '第三篇：溝通與協作（職場對話術）'
        if num <= 14: return '第四篇：職場生存與防線（自我保護與邊界）'
        if num <= 17: return '第五篇：職涯與個人成長（價值提升與升遷）'
        return '第六篇：能量、韌性與生活平衡（心智與身心管理）'

    current_chapter = None
    
    for p_item in principles:
        num = p_item['num']
        chap_name = get_chapter_name(num)
        
        # Add Chapter Heading if changed
        if chap_name != current_chapter:
            current_chapter = chap_name
            h_chap = H(outlinelevel=1, stylename=chapter_style)
            h_chap.addText(current_chapter)
            doc.text.addElement(h_chap)
            
            if chapters:
                chap_id = '1' if num <= 4 else ('2' if num <= 8 else ('3' if num <= 11 else ('4' if num <= 14 else ('5' if num <= 17 else '6'))))
                matched_chap = next((c for c in chapters if c['id'] == chap_id), None)
                if matched_chap and matched_chap.get('quotes'):
                    for q in matched_chap['quotes']:
                        p_q = P(stylename=principle_body_style)
                        p_q.addText(f"• {q}")
                        doc.text.addElement(p_q)
            
        # Add Principle Title (Always starts on new page)
        h_pr = H(outlinelevel=2, stylename=principle_title_style)
        h_pr.addText(f"原則 {num}：{p_item['title']}")
        doc.text.addElement(h_pr)
        
        # Add Principle Content
        def add_translation_row(doc, label, text):
            p = P(stylename=principle_body_style)
            span = Span(stylename=bold_label_style)
            span.addText(label + "：")
            p.addElement(span)
            p.addText(text)
            doc.text.addElement(p)
            
        add_translation_row(doc, "國語 (Mandarin)", p_item['mandarin'])
        add_translation_row(doc, "English", p_item['english'])
        add_translation_row(doc, "台語 (Taiwanese)", p_item['taiwanese'])
        
        # Add Slide Items corresponding to this Principle
        refs = [r.strip() for r in p_item['ref'].split(',')]
        matched_items = []
        for item in final_items:
            item_key = item['key']
            base_key = item_key.split('-')[0]
            if item_key in refs or base_key in refs:
                matched_items.append(item)
                
        if matched_items:
            for idx, item in enumerate(matched_items):
                # Page break before first item (so Principle is on separate page)
                # and after every 2 items (so max 2 items per page)
                h_style = slide_title_style_with_break if idx % 2 == 0 else slide_title_style
                
                # Add Slide Item Title
                h_slide = H(outlinelevel=3, stylename=h_style)
                h_slide.addText(f"No. {item['key']} (Slide {item['slide']})")
                doc.text.addElement(h_slide)
                
                # Original Text
                p_orig = P(stylename=slide_orig_style)
                span = Span(stylename=bold_label_style)
                span.addText("原文：")
                p_orig.addElement(span)
                p_orig.addText(item['original'])
                doc.text.addElement(p_orig)
                
                # Mandarin
                p_mand = P(stylename=slide_trans_style)
                span = Span(stylename=bold_label_style)
                span.addText("優化中文：")
                p_mand.addElement(span)
                p_mand.addText(item['mandarin'])
                doc.text.addElement(p_mand)
                
                # English
                p_eng = P(stylename=slide_trans_style)
                span = Span(stylename=bold_label_style)
                span.addText("優化英文：")
                p_eng.addElement(span)
                p_eng.addText(item['english'])
                doc.text.addElement(p_eng)
                
                # Taiwanese
                p_taiw = P(stylename=slide_trans_style)
                span = Span(stylename=bold_label_style)
                span.addText("優化台文：")
                p_taiw.addElement(span)
                p_taiw.addText(item['taiwanese'])
                doc.text.addElement(p_taiw)
                
    doc.save(odt_path)
    print(f"Saved regenerated ODT handbook: {odt_path}")

def main():
    pptx_path = "工作原則-書.pptx.pptx"
    md_path = "工作原則整理與潤稿.md"
    db_path = "translation_db.json"
    
    print("Step 1: Parsing PowerPoint slides...")
    pptx_items = parse_pptx_items(pptx_path)
    print(f"Found {len(pptx_items)} items in PPTX.")
    
    print("\nStep 2: Loading translation database and manual overrides...")
    db = {}
    if os.path.exists(db_path):
        with open(db_path, "r", encoding="utf-8-sig") as f:
            db = json.load(f)
    print(f"Loaded {len(db)} entries from database.")
    
    overrides_path = "manual_overrides.json"
    overrides = {}
    if os.path.exists(overrides_path):
        try:
            with open(overrides_path, "r", encoding="utf-8-sig") as f:
                raw_overrides = json.load(f)
                overrides = {k: v for k, v in raw_overrides.items() if not k.startswith("_")}
            if overrides:
                print(f"Loaded {len(overrides)} active manual override rules from {overrides_path}.")
        except Exception as e:
            print(f"Warning: Could not parse {overrides_path}: {e}")
    
    # Track missing or modified items to translate
    missing_items = []
    
    print("\nStep 3: Matching items and verifying translations...")
    final_items = []
    translated_new_count = 0
    
    for item in pptx_items:
        key = item['key']
        original = item['text']
        
        # Check manual overrides first (by key, e.g. "12", or by original text)
        override_match = overrides.get(str(key)) or overrides.get(original)
        
        # Check database by original text match
        db_match = db.get(original)
        
        # If not found, try to look up with minor space differences
        if not db_match:
            original_clean = re.sub(r'\s+', '', original)
            for db_orig, db_val in db.items():
                if re.sub(r'\s+', '', db_orig) == original_clean:
                    db_match = db_val
                    break
                    
        if db_match or override_match:
            base_eng = db_match['english'] if db_match else original
            base_man = db_match['mandarin'] if db_match else ""
            base_tai = db_match['taiwanese'] if db_match else ""
            
            if override_match and isinstance(override_match, dict):
                eng = override_match.get('english', base_eng)
                man = override_match.get('mandarin', base_man)
                tai = override_match.get('taiwanese', base_tai)
            else:
                eng, man, tai = base_eng, base_man, base_tai

            final_items.append({
                "key": key,
                "slide": item['slide'],
                "original": original,
                "english": eng,
                "mandarin": man,
                "taiwanese": tai
            })
        else:
            # We need to translate this
            print(f"  Missing translation for Key {key}: '{original[:60]}...'")
            missing_items.append((original, key, item['slide']))

    # Perform translations if API key is present
    api_key = os.environ.get("GEMINI_API_KEY")
    if missing_items:
        if api_key:
            print(f"\nFound {len(missing_items)} items requiring translation. Running API translations...")
            for idx, (orig, key, slide) in enumerate(missing_items, 1):
                print(f"  ({idx}/{len(missing_items)}) Translating Key {key}...")
                translated = translate_via_gemini(orig)
                if translated:
                    # Save to db
                    db[orig] = {
                        "english": translated['english'],
                        "mandarin": translated['mandarin'],
                        "taiwanese": translated['taiwanese'],
                        "key": key
                    }
                    final_items.append({
                        "key": key,
                        "slide": slide,
                        "original": orig,
                        "english": translated['english'],
                        "mandarin": translated['mandarin'],
                        "taiwanese": translated['taiwanese']
                    })
                    translated_new_count += 1
                else:
                    # Fallback placeholders
                    final_items.append({
                        "key": key,
                        "slide": slide,
                        "original": orig,
                        "english": f"[TBD] {orig}",
                        "mandarin": f"[TBD] 待翻譯 - 原文: {orig}",
                        "taiwanese": f"[TBD] 待翻譯 - 原文: {orig}"
                    })
            # Save updated DB
            with open(db_path, "w", encoding="utf-8") as f:
                json.dump(db, f, ensure_ascii=False, indent=2)
            print(f"API translation complete. Added {translated_new_count} translations to database.")
        else:
            print(f"\n[WARNING] Found {len(missing_items)} items requiring translation, but GEMINI_API_KEY is not set.")
            print("Placeholders will be generated. Please set GEMINI_API_KEY in `.env` to automatically translate them.")
            for orig, key, slide in missing_items:
                final_items.append({
                    "key": key,
                    "slide": slide,
                    "original": orig,
                    "english": f"[TBD] {orig}",
                    "mandarin": f"[TBD] 待翻譯 - 原文: {orig}",
                    "taiwanese": f"[TBD] 待翻譯 - 原文: {orig}"
                })

    # Sort final items by key number
    def get_key_sort(x):
        k = x['key']
        # Split number and suffix
        parts = re.split(r'(\d+)', k)
        return [int(s) if s.isdigit() else s for s in parts]
        
    final_items.sort(key=get_key_sort)
    
    print("\nStep 4: Parsing Part 1 (Principles) from existing markdown...")
    part1_content, principles, chapters = parse_existing_md(md_path)
    
    # Apply manual overrides to core principles if specified
    for p in principles:
        p_num = p['num']
        p_override = (
            overrides.get(f"principle_{p_num}")
            or overrides.get(f"P{p_num}")
            or overrides.get(f"原則_{p_num}")
            or overrides.get(f"原則{p_num}")
        )
        if p_override and isinstance(p_override, dict):
            if 'title' in p_override: p['title'] = p_override['title']
            if 'mandarin' in p_override: p['mandarin'] = p_override['mandarin']
            if 'english' in p_override: p['english'] = p_override['english']
            if 'taiwanese' in p_override: p['taiwanese'] = p_override['taiwanese']
    
    if not part1_content:
        # Fallback default title/toc if file does not exist
        part1_content = "# 工作的管見：職場生存與成長的避坑指南\n## 工作原則整理與潤稿 (三語對照版)\n\n"
        
    print("\nStep 5: Regenerating Markdown handbook...")
    # Write Part 1 and regenerated Part 2
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(part1_content.strip() + "\n\n---\n\n## 第二部分：簡報對照與逐條潤稿 (簡報版)\n\n")
        f.write(f"> [!NOTE]\n> 本部分為逐頁對照簡報 (PPTX) 的 {len(final_items)} 條原始筆記進行優化與翻譯的內容，方便您在更新投影片時直接複製貼上使用。\n\n")
        
        current_slide = 0
        for item in final_items:
            slide = item['slide']
            if slide != current_slide:
                current_slide = slide
                # Find min/max keys on this slide
                slide_items = [x for x in final_items if x['slide'] == slide]
                min_key = slide_items[0]['key']
                max_key = slide_items[-1]['key']
                f.write(f"### Slide {slide} (原則 {min_key} - {max_key})\n\n")
                
            # Find matched principles
            matched_principles = []
            item_key = item['key']
            base_key = item_key.split('-')[0]
            for p in principles:
                refs = [r.strip() for r in p['ref'].split(',')]
                if item_key in refs or base_key in refs:
                    matched_principles.append(p['num'])
            
            if matched_principles:
                principles_str = ", ".join(f"原則 {num}" for num in matched_principles)
                f.write(f"#### **No. {item['key']}** ({principles_str})\n")
            else:
                f.write(f"#### **No. {item['key']}**\n")
            f.write(f"*   **原文**：`{item['original']}`\n")
            f.write(f"*   **優化英文**：{item['english']}\n")
            f.write(f"*   **優化中文**：{item['mandarin']}\n")
            f.write(f"*   **優化台文**：{item['taiwanese']}\n\n")
            
    print(f"Saved regenerated handbook: {md_path}")
    
    print("\nStep 6: Generating HTML card deck web application...")
    
    # We serialize slides and principles as JSON variables in HTML
    principles_json = json.dumps(principles, ensure_ascii=False)
    slides_json = json.dumps(final_items, ensure_ascii=False)
    chapters_json = json.dumps(chapters, ensure_ascii=False)

    slide_groups = {}
    for item in final_items:
        s = item['slide']
        if s not in slide_groups: slide_groups[s] = []
        slide_groups[s].append(item['key'])
    slide_buttons = ['<button class="filter-tag active" onclick="filterSlides(\'all\')">全部 Slide</button>']
    for s in sorted(slide_groups.keys()):
        min_k = slide_groups[s][0]
        max_k = slide_groups[s][-1]
        slide_buttons.append(f'<button class="filter-tag" onclick="filterSlides(\'{s}\')">Slide {s} ({min_k}-{max_k})</button>')
    slides_filter_html = "\n                ".join(slide_buttons)
    
    build_timestamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    
    # HTML template with embedded styling and logic
    html_content = f"""<!DOCTYPE html>
<html lang="zh-Hant">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>工作的管見 (My Two Cents) - 工作原則卡牌</title>
    <meta name="description" content="工作的管見 (My Two Cents)：From the Floor: My Two Cents on Practical Work Principles - 職場生存與成長的避坑指南">
    <link href="https://fonts.googleapis.com/css2?family=Outfit:wght@300;400;600;800&family=Noto+Sans+TC:wght@300;400;700&display=swap" rel="stylesheet">
    <style>
        :root {{
            --primary: #6366f1;
            --primary-hover: #4f46e5;
            --bg-dark: #0f172a;
            --bg-card: rgba(30, 41, 59, 0.7);
            --bg-glass: rgba(15, 23, 42, 0.65);
            --text-main: #f8fafc;
            --text-muted: #94a3b8;
            --border: rgba(255, 255, 255, 0.08);
            --border-hover: rgba(99, 102, 241, 0.4);
            --glow: rgba(99, 102, 241, 0.15);
        }}

        .light-mode {{
            --primary: #4f46e5;
            --primary-hover: #3730a3;
            --bg-dark: #f8fafc;
            --bg-card: rgba(255, 255, 255, 0.85);
            --bg-glass: rgba(241, 245, 249, 0.75);
            --text-main: #0f172a;
            --text-muted: #64748b;
            --border: rgba(0, 0, 0, 0.08);
            --border-hover: rgba(79, 70, 229, 0.4);
            --glow: rgba(79, 70, 229, 0.1);
        }}

        * {{
            box-sizing: border-box;
            margin: 0;
            padding: 0;
            transition: background-color 0.3s, border-color 0.3s;
        }}

        body {{
            font-family: 'Outfit', 'Noto Sans TC', sans-serif;
            background-color: var(--bg-dark);
            color: var(--text-main);
            min-height: 100vh;
            display: flex;
            flex-direction: column;
            overflow-x: hidden;
            background-image: 
                radial-gradient(at 0% 0%, rgba(99, 102, 241, 0.12) 0px, transparent 50%),
                radial-gradient(at 100% 100%, rgba(139, 92, 246, 0.1) 0px, transparent 50%);
            background-attachment: fixed;
        }}

        /* Header block */
        header {{
            padding: 2.5rem 1.5rem 1.5rem;
            text-align: center;
            border-bottom: 1px solid var(--border);
            background: var(--bg-glass);
            backdrop-filter: blur(20px);
            -webkit-backdrop-filter: blur(20px);
            position: sticky;
            top: 0;
            z-index: 100;
        }}

        header h1 {{
            font-size: 2.2rem;
            font-weight: 800;
            letter-spacing: -0.025em;
            background: linear-gradient(135deg, #a5b4fc 0%, #6366f1 50%, #818cf8 100%);
            -webkit-background-clip: text;
            -webkit-text-fill-color: transparent;
            margin-bottom: 0.35rem;
            display: inline-flex;
            align-items: baseline;
            gap: 0.5rem;
            flex-wrap: wrap;
        }}

        .light-mode header h1 {{
            background: linear-gradient(135deg, #4f46e5 0%, #312e81 100%);
            -webkit-background-clip: text;
            -webkit-text-fill-color: transparent;
        }}

        .en-title {{
            font-size: 1.35rem;
            font-weight: 600;
            font-family: 'Outfit', sans-serif;
            letter-spacing: 0.01em;
            opacity: 0.9;
        }}

        header p {{
            color: var(--text-muted);
            font-size: 0.95rem;
            font-weight: 400;
            line-height: 1.5;
            text-align: left;
        }}

        .header-sub-en {{
            display: block;
            font-size: 0.85rem;
            color: var(--primary);
            font-weight: 500;
            font-family: 'Outfit', sans-serif;
            letter-spacing: 0.01em;
            margin-top: 0.2rem;
            opacity: 0.95;
        }}

        .light-mode .header-sub-en {{
            color: #4f46e5;
        }}

        /* Navigation and tools */
        .controls-container {{
            max-width: 1200px;
            width: 100%;
            margin: 0 auto;
            display: flex;
            flex-wrap: wrap;
            gap: 1rem;
            align-items: center;
            justify-content: space-between;
        }}

        .tabs {{
            display: flex;
            background: rgba(0, 0, 0, 0.2);
            padding: 0.25rem;
            border-radius: 99px;
            border: 1px solid var(--border);
        }}

        .light-mode .tabs {{
            background: rgba(0, 0, 0, 0.05);
        }}

        .tab-btn {{
            background: transparent;
            border: none;
            color: var(--text-muted);
            padding: 0.6rem 1.4rem;
            font-size: 0.95rem;
            font-weight: 600;
            border-radius: 99px;
            cursor: pointer;
            transition: all 0.2s cubic-bezier(0.4, 0, 0.2, 1);
        }}

        .tab-btn.active {{
            background: var(--primary);
            color: #ffffff;
            box-shadow: 0 4px 12px var(--glow);
        }}

        .action-tools {{
            display: flex;
            align-items: center;
            gap: 0.75rem;
        }}

        .search-bar {{
            position: relative;
            min-width: 260px;
        }}

        .search-input {{
            width: 100%;
            padding: 0.6rem 1rem 0.6rem 2.5rem;
            background: rgba(0, 0, 0, 0.2);
            border: 1px solid var(--border);
            border-radius: 99px;
            color: var(--text-main);
            font-size: 0.9rem;
            outline: none;
            transition: all 0.2s;
        }}

        .light-mode .search-input {{
            background: rgba(255, 255, 255, 0.9);
        }}

        .search-input:focus {{
            border-color: var(--primary);
            box-shadow: 0 0 0 3px var(--glow);
        }}

        .search-icon {{
            position: absolute;
            left: 1rem;
            top: 50%;
            transform: translateY(-50%);
            color: var(--text-muted);
            pointer-events: none;
            font-size: 0.9rem;
        }}

        .icon-btn {{
            background: rgba(0, 0, 0, 0.2);
            border: 1px solid var(--border);
            width: 40px;
            height: 40px;
            border-radius: 50%;
            display: flex;
            align-items: center;
            justify-content: center;
            cursor: pointer;
            color: var(--text-main);
            outline: none;
            font-size: 1.1rem;
        }}

        .light-mode .icon-btn {{
            background: rgba(255, 255, 255, 0.9);
        }}

        .icon-btn:hover {{
            border-color: var(--primary);
            color: var(--primary);
        }}

        /* Main layouts */
        main {{
            flex: 1;
            max-width: 1200px;
            width: 100%;
            margin: 2rem auto;
            padding: 0 1.5rem;
        }}

        .section-view {{
            display: none;
        }}

        .section-view.active {{
            display: block;
            animation: fadeIn 0.4s ease-out;
        }}

        @keyframes fadeIn {{
            from {{ opacity: 0; transform: translateY(10px); }}
            to {{ opacity: 1; transform: translateY(0); }}
        }}

        /* Filters */
        .filter-tags {{
            display: flex;
            flex-wrap: wrap;
            gap: 0.5rem;
            margin-bottom: 1.5rem;
        }}

        .filter-tag {{
            background: rgba(255, 255, 255, 0.03);
            border: 1px solid var(--border);
            padding: 0.4rem 1rem;
            border-radius: 99px;
            font-size: 0.85rem;
            cursor: pointer;
            color: var(--text-muted);
            font-weight: 500;
        }}

        .light-mode .filter-tag {{
            background: rgba(0, 0, 0, 0.02);
        }}

        .filter-tag:hover, .filter-tag.active {{
            border-color: var(--primary);
            color: var(--primary);
            background: rgba(99, 102, 241, 0.05);
        }}

        /* Card grid */
        .card-grid {{
            display: grid;
            grid-template-columns: repeat(auto-fill, minmax(340px, 1fr));
            gap: 1.5rem;
        }}

        /* Chapter Header Banner with Golden Quotes */
        .chapter-header-banner {{
            grid-column: 1 / -1;
            background: linear-gradient(135deg, rgba(30, 41, 59, 0.75) 0%, rgba(15, 23, 42, 0.85) 100%);
            border: 1px solid rgba(99, 102, 241, 0.25);
            border-radius: 16px;
            padding: 1.4rem 1.6rem;
            margin-top: 1.5rem;
            margin-bottom: 0.25rem;
            backdrop-filter: blur(12px);
            -webkit-backdrop-filter: blur(12px);
            box-shadow: 0 4px 20px rgba(0, 0, 0, 0.15);
            transition: all 0.3s ease;
        }}

        .chapter-header-banner:first-child {{
            margin-top: 0;
        }}

        .light-mode .chapter-header-banner {{
            background: linear-gradient(135deg, rgba(248, 250, 252, 0.95) 0%, rgba(241, 245, 249, 0.9) 100%);
            border-color: rgba(99, 102, 241, 0.2);
            box-shadow: 0 4px 20px rgba(99, 102, 241, 0.06);
        }}

        .chapter-banner-top {{
            display: flex;
            align-items: center;
            gap: 0.8rem;
            margin-bottom: 1rem;
            flex-wrap: wrap;
        }}

        .chapter-badge {{
            font-size: 0.75rem;
            font-weight: 800;
            letter-spacing: 0.06em;
            text-transform: uppercase;
            color: #ffffff;
            background: linear-gradient(135deg, #4f46e5 0%, #7c3aed 100%);
            padding: 0.3rem 0.7rem;
            border-radius: 6px;
            box-shadow: 0 2px 8px rgba(79, 70, 229, 0.3);
        }}

        .chapter-title {{
            font-size: 1.25rem;
            font-weight: 700;
            color: var(--text-main);
            margin: 0;
            letter-spacing: -0.01em;
        }}

        .chapter-range-chip {{
            font-size: 0.8rem;
            font-weight: 600;
            color: var(--primary);
            background: rgba(99, 102, 241, 0.1);
            padding: 0.25rem 0.65rem;
            border-radius: 99px;
            border: 1px solid rgba(99, 102, 241, 0.2);
        }}

        .light-mode .chapter-range-chip {{
            background: rgba(99, 102, 241, 0.08);
        }}

        .chapter-quotes-area {{
            background: rgba(15, 23, 42, 0.4);
            border: 1px dashed rgba(99, 102, 241, 0.2);
            border-radius: 12px;
            padding: 0.9rem 1.1rem;
        }}

        .light-mode .chapter-quotes-area {{
            background: rgba(255, 255, 255, 0.75);
            border-color: rgba(99, 102, 241, 0.2);
        }}

        .chapter-quotes-header {{
            font-size: 0.8rem;
            font-weight: 700;
            color: #818cf8;
            margin-bottom: 0.6rem;
            display: flex;
            align-items: center;
            gap: 0.4rem;
        }}

        .light-mode .chapter-quotes-header {{
            color: #4f46e5;
        }}

        .chapter-quotes-pills {{
            display: flex;
            flex-wrap: wrap;
            gap: 0.5rem;
        }}

        .quote-pill {{
            display: inline-flex;
            align-items: center;
            font-size: 0.82rem;
            font-weight: 600;
            color: var(--text-main);
            background: rgba(99, 102, 241, 0.08);
            border: 1px solid rgba(99, 102, 241, 0.2);
            border-radius: 99px;
            padding: 0.35rem 0.8rem;
            letter-spacing: 0.02em;
            transition: all 0.2s cubic-bezier(0.4, 0, 0.2, 1);
            cursor: default;
            user-select: none;
        }}

        .light-mode .quote-pill {{
            background: #ffffff;
            border-color: rgba(99, 102, 241, 0.2);
            color: #334155;
            box-shadow: 0 1px 3px rgba(0, 0, 0, 0.04);
        }}

        .quote-pill:hover {{
            border-color: var(--primary);
            background: rgba(99, 102, 241, 0.18);
            transform: translateY(-2px);
            box-shadow: 0 4px 12px rgba(99, 102, 241, 0.15);
            color: var(--primary);
        }}

        .light-mode .quote-pill:hover {{
            background: rgba(99, 102, 241, 0.08);
            color: var(--primary);
        }}

        .principle-card {{
            background: var(--bg-card);
            border: 1px solid var(--border);
            border-radius: 16px;
            padding: 1.5rem;
            backdrop-filter: blur(10px);
            -webkit-backdrop-filter: blur(10px);
            box-shadow: 0 4px 30px rgba(0, 0, 0, 0.1);
            display: flex;
            flex-direction: column;
            position: relative;
            overflow: hidden;
            cursor: pointer;
            transition: all 0.3s cubic-bezier(0.4, 0, 0.2, 1);
        }}

        .principle-card::before {{
            content: '';
            position: absolute;
            top: 0;
            left: 0;
            right: 0;
            height: 4px;
            background: linear-gradient(90deg, var(--primary) 0%, #818cf8 100%);
            opacity: 0;
            transition: opacity 0.3s;
        }}

        .principle-card:hover {{
            transform: translateY(-5px);
            border-color: var(--border-hover);
            box-shadow: 0 12px 30px rgba(99, 102, 241, 0.08);
        }}

        .principle-card:hover::before {{
            opacity: 1;
        }}

        .card-header {{
            display: flex;
            justify-content: space-between;
            align-items: center;
            margin-bottom: 1rem;
        }}

        .card-num {{
            font-size: 0.8rem;
            font-weight: 800;
            text-transform: uppercase;
            letter-spacing: 0.05em;
            color: var(--primary);
            background: rgba(99, 102, 241, 0.1);
            padding: 0.25rem 0.6rem;
            border-radius: 4px;
        }}

        .card-principle-badge {{
            font-size: 0.8rem;
            font-weight: 800;
            letter-spacing: 0.05em;
            color: #10b981;
            background: rgba(16, 185, 129, 0.1);
            padding: 0.25rem 0.6rem;
            border-radius: 4px;
            display: inline-block;
        }}

        .light-mode .card-principle-badge {{
            color: #059669;
            background: rgba(5, 150, 105, 0.1);
        }}

        .ref-items-container {{
            margin-top: 1.2rem;
            padding-top: 0.8rem;
            border-top: 1px dashed var(--border);
        }}

        .ref-items-label {{
            font-size: 0.78rem;
            font-weight: 600;
            color: var(--text-muted);
            margin-bottom: 0.6rem;
            display: block;
        }}

        .ref-chips-grid {{
            display: flex;
            flex-wrap: wrap;
            gap: 0.4rem;
        }}

        .ref-item-chip {{
            font-size: 0.75rem;
            font-weight: 600;
            font-family: inherit;
            color: var(--primary);
            background: rgba(99, 102, 241, 0.12);
            border: 1px solid rgba(99, 102, 241, 0.25);
            padding: 0.25rem 0.55rem;
            border-radius: 6px;
            cursor: pointer;
            transition: all 0.2s ease;
        }}

        .light-mode .ref-item-chip {{
            background: rgba(79, 70, 229, 0.08);
            border-color: rgba(79, 70, 229, 0.2);
            color: var(--primary);
        }}

        .ref-item-chip:hover {{
            background: var(--primary);
            color: #ffffff;
            border-color: var(--primary);
            transform: translateY(-2px);
            box-shadow: 0 4px 12px rgba(99, 102, 241, 0.35);
        }}

        .highlight-pulse {{
            animation: pulseGlow 2.5s ease-in-out infinite;
            border-color: var(--primary) !important;
            z-index: 10;
        }}

        @keyframes pulseGlow {{
            0% {{ box-shadow: 0 0 0 0 rgba(99, 102, 241, 0.8); transform: scale(1); }}
            50% {{ box-shadow: 0 0 30px 10px rgba(99, 102, 241, 0.6); transform: scale(1.03); }}
            100% {{ box-shadow: 0 0 0 0 rgba(99, 102, 241, 0); transform: scale(1); }}
        }}

        .card-meta {{
            font-size: 0.8rem;
            color: var(--text-muted);
        }}

        .card-title {{
            font-size: 1.2rem;
            font-weight: 700;
            margin-bottom: 1rem;
            line-height: 1.3;
        }}

        .lang-section {{
            margin-bottom: 0.75rem;
            font-size: 0.95rem;
            line-height: 1.5;
        }}

        .lang-section:last-child {{
            margin-bottom: 0;
        }}

        .lang-label {{
            font-size: 0.75rem;
            font-weight: 800;
            color: var(--text-muted);
            text-transform: uppercase;
            letter-spacing: 0.05em;
            display: block;
            margin-bottom: 0.2rem;
        }}

        .lang-txt {{
            color: var(--text-main);
        }}

        .original-eng {{
            font-style: italic;
            color: var(--text-muted);
            border-left: 2px solid var(--border);
            padding-left: 0.75rem;
            margin-bottom: 1rem;
            font-size: 0.9rem;
        }}

        /* Random Card drawer view */
        .random-view-container {{
            max-width: 580px;
            margin: 2rem auto;
            display: flex;
            flex-direction: column;
            align-items: center;
        }}

        /* 3D card flip setup */
        .flip-card-container {{
            perspective: 1200px;
            width: 100%;
            height: 480px;
            cursor: pointer;
            margin-bottom: 2rem;
        }}

        .flip-card {{
            position: relative;
            width: 100%;
            height: 100%;
            transition: transform 0.6s cubic-bezier(0.4, 0, 0.2, 1);
            transform-style: preserve-3d;
        }}

        .flip-card-container.flipped .flip-card {{
            transform: rotateY(180deg);
        }}

        .card-face {{
            position: absolute;
            width: 100%;
            height: 100%;
            backface-visibility: hidden;
            -webkit-backface-visibility: hidden;
            border-radius: 24px;
            border: 1px solid var(--border);
            padding: 2rem;
            box-shadow: 0 15px 35px rgba(0,0,0,0.2);
            display: flex;
            flex-direction: column;
            justify-content: space-between;
            background: var(--bg-card);
            backdrop-filter: blur(25px);
            -webkit-backdrop-filter: blur(25px);
        }}

        .card-face-front {{
            background: linear-gradient(135deg, rgba(30, 41, 59, 0.8) 0%, rgba(15, 23, 42, 0.9) 100%);
            border-color: rgba(99, 102, 241, 0.25);
        }}

        .light-mode .card-face-front {{
            background: linear-gradient(135deg, rgba(255, 255, 255, 0.9) 0%, rgba(241, 245, 249, 0.95) 100%);
        }}

        .card-face-back {{
            transform: rotateY(180deg);
            background: linear-gradient(135deg, rgba(15, 23, 42, 0.95) 0%, rgba(30, 41, 59, 0.8) 100%);
            border-color: rgba(139, 92, 246, 0.25);
        }}

        .light-mode .card-face-back {{
            background: linear-gradient(135deg, rgba(241, 245, 249, 0.95) 0%, rgba(255, 255, 255, 0.9) 100%);
        }}

        .front-body {{
            display: flex;
            flex-direction: column;
            justify-content: center;
            align-items: center;
            flex: 1;
            text-align: center;
        }}

        .front-title {{
            font-size: 1.5rem;
            font-weight: 600;
            line-height: 1.5;
            color: var(--text-main);
        }}

        .back-body {{
            display: flex;
            flex-direction: column;
            justify-content: center;
            gap: 1.25rem;
            flex: 1;
        }}

        .flip-hint {{
            font-size: 0.8rem;
            color: var(--text-muted);
            text-align: center;
            letter-spacing: 0.05em;
            text-transform: uppercase;
        }}

        .draw-btn {{
            background: linear-gradient(90deg, var(--primary) 0%, #818cf8 100%);
            color: white;
            border: none;
            padding: 0.9rem 2.5rem;
            font-size: 1.05rem;
            font-weight: 700;
            border-radius: 99px;
            cursor: pointer;
            box-shadow: 0 4px 20px var(--glow);
            transition: all 0.2s cubic-bezier(0.4, 0, 0.2, 1);
            display: flex;
            align-items: center;
            gap: 0.5rem;
        }}

        .draw-btn:hover {{
            transform: scale(1.05);
            box-shadow: 0 8px 25px var(--glow);
        }}

        .draw-btn:active {{
            transform: scale(0.98);
        }}

        /* Particle effect */
        .sparkle-particles {{
            position: absolute;
            pointer-events: none;
            width: 10px;
            height: 10px;
            background: var(--primary);
            border-radius: 50%;
            opacity: 0;
        }}

        /* VIEW 4: Connection Graph (關係圖譜) */
        .graph-wrapper {{
            display: flex;
            flex-direction: column;
            width: 100%;
            background: var(--bg-card);
            border: 1px solid var(--border);
            border-radius: 18px;
            backdrop-filter: blur(12px);
            -webkit-backdrop-filter: blur(12px);
            overflow: hidden;
            box-shadow: 0 8px 32px rgba(0, 0, 0, 0.2);
            position: relative;
        }}

        .graph-toolbar {{
            display: flex;
            flex-wrap: wrap;
            justify-content: space-between;
            align-items: center;
            gap: 0.75rem;
            padding: 0.85rem 1.25rem;
            background: rgba(0, 0, 0, 0.18);
            border-bottom: 1px solid var(--border);
        }}

        .light-mode .graph-toolbar {{
            background: rgba(0, 0, 0, 0.03);
        }}

        .graph-toolbar-left {{
            display: flex;
            flex-wrap: wrap;
            align-items: center;
            gap: 0.85rem;
        }}

        .graph-title {{
            font-size: 0.95rem;
            font-weight: 700;
            color: var(--text-main);
            display: flex;
            align-items: center;
            gap: 0.4rem;
        }}

        .graph-filters {{
            display: flex;
            flex-wrap: wrap;
            gap: 0.35rem;
            align-items: center;
        }}

        .graph-mode-switch {{
            display: inline-flex;
            background: rgba(0, 0, 0, 0.28);
            border: 1px solid var(--border);
            border-radius: 99px;
            padding: 2px;
            gap: 2px;
        }}

        .light-mode .graph-mode-switch {{
            background: rgba(0, 0, 0, 0.05);
        }}

        .graph-mode-btn {{
            background: transparent;
            border: none;
            color: var(--text-muted);
            font-size: 0.76rem;
            font-weight: 600;
            padding: 0.25rem 0.65rem;
            border-radius: 99px;
            cursor: pointer;
            transition: all 0.2s;
            display: inline-flex;
            align-items: center;
            gap: 0.25rem;
            font-family: inherit;
        }}

        .graph-mode-btn.active {{
            background: var(--primary);
            color: #ffffff;
            box-shadow: 0 2px 8px rgba(99, 102, 241, 0.35);
        }}

        .graph-mode-btn:not(.active):hover {{
            color: var(--text-main);
        }}

        .drawer-action-btn {{
            background: rgba(99, 102, 241, 0.12);
            border: 1px solid rgba(99, 102, 241, 0.3);
            color: var(--primary);
            font-size: 0.74rem;
            font-weight: 600;
            padding: 0.22rem 0.65rem;
            border-radius: 99px;
            cursor: pointer;
            transition: all 0.2s;
            display: inline-flex;
            align-items: center;
            gap: 0.3rem;
            font-family: inherit;
        }}

        .drawer-action-btn:hover {{
            background: var(--primary);
            color: #ffffff;
        }}

        .drawer-action-btn.active {{
            background: #10b981;
            border-color: #10b981;
            color: #ffffff;
        }}

        .graph-toolbar-right {{
            display: flex;
            align-items: center;
            gap: 0.5rem;
        }}

        .graph-btn {{
            background: rgba(99, 102, 241, 0.12);
            border: 1px solid rgba(99, 102, 241, 0.25);
            color: var(--primary);
            font-size: 0.82rem;
            font-weight: 600;
            padding: 0.35rem 0.75rem;
            border-radius: 8px;
            cursor: pointer;
            transition: all 0.2s ease;
            display: inline-flex;
            align-items: center;
            gap: 0.3rem;
            font-family: inherit;
        }}

        .light-mode .graph-btn {{
            background: rgba(79, 70, 229, 0.08);
            border-color: rgba(79, 70, 229, 0.2);
            color: var(--primary);
        }}

        .graph-btn:hover {{
            background: var(--primary);
            color: #ffffff;
            border-color: var(--primary);
            transform: translateY(-1px);
        }}

        .graph-canvas-container {{
            position: relative;
            width: 100%;
            height: 720px;
            overflow: hidden;
            background: radial-gradient(circle at center, rgba(30, 41, 59, 0.4) 0%, rgba(15, 23, 42, 0.8) 100%);
            cursor: grab;
        }}

        .light-mode .graph-canvas-container {{
            background: radial-gradient(circle at center, rgba(241, 245, 249, 0.8) 0%, rgba(226, 232, 240, 0.5) 100%);
        }}

        .graph-canvas-container.panning {{
            cursor: grabbing;
        }}

        #graphCanvas {{
            display: block;
            width: 100%;
            height: 100%;
        }}

        .graph-legend {{
            position: absolute;
            top: 1rem;
            left: 1rem;
            display: flex;
            flex-wrap: wrap;
            gap: 0.5rem 0.8rem;
            background: rgba(15, 23, 42, 0.82);
            border: 1px solid var(--border);
            border-radius: 10px;
            padding: 0.4rem 0.85rem;
            font-size: 0.76rem;
            font-weight: 500;
            color: var(--text-muted);
            backdrop-filter: blur(8px);
            -webkit-backdrop-filter: blur(8px);
            pointer-events: none;
            z-index: 5;
            box-shadow: 0 4px 15px rgba(0, 0, 0, 0.2);
        }}

        .light-mode .graph-legend {{
            background: rgba(255, 255, 255, 0.88);
            color: var(--text-main);
        }}

        .legend-chip {{
            display: inline-flex;
            align-items: center;
            gap: 0.35rem;
        }}

        .legend-dot {{
            width: 8px;
            height: 8px;
            border-radius: 50%;
        }}

        .graph-hint {{
            position: absolute;
            bottom: 1rem;
            left: 1rem;
            background: rgba(15, 23, 42, 0.8);
            border: 1px solid var(--border);
            border-radius: 8px;
            padding: 0.35rem 0.75rem;
            font-size: 0.74rem;
            color: var(--text-muted);
            backdrop-filter: blur(8px);
            -webkit-backdrop-filter: blur(8px);
            pointer-events: none;
            z-index: 5;
        }}

        .light-mode .graph-hint {{
            background: rgba(255, 255, 255, 0.85);
            color: #64748b;
        }}

        /* Sliding Detail Drawer */
        .graph-drawer {{
            position: absolute;
            top: 0;
            right: 0;
            bottom: 0;
            width: 440px;
            max-width: 92%;
            background: rgba(15, 23, 42, 0.95);
            backdrop-filter: blur(16px);
            -webkit-backdrop-filter: blur(16px);
            border-left: 1px solid var(--border);
            z-index: 20;
            display: flex;
            flex-direction: column;
            transform: translateX(100%);
            transition: transform 0.35s cubic-bezier(0.16, 1, 0.3, 1);
            box-shadow: -8px 0 35px rgba(0, 0, 0, 0.35);
        }}

        .light-mode .graph-drawer {{
            background: rgba(255, 255, 255, 0.96);
        }}

        .graph-drawer.open {{
            transform: translateX(0);
        }}

        .drawer-header {{
            display: flex;
            align-items: center;
            justify-content: space-between;
            padding: 1.1rem 1.4rem;
            border-bottom: 1px solid var(--border);
            background: rgba(0, 0, 0, 0.15);
        }}

        .light-mode .drawer-header {{
            background: rgba(0, 0, 0, 0.02);
        }}

        .drawer-title-group {{
            display: flex;
            align-items: center;
            gap: 0.5rem;
            flex-wrap: wrap;
        }}

        .drawer-badge {{
            font-size: 0.8rem;
            font-weight: 800;
            color: var(--primary);
            background: rgba(99, 102, 241, 0.12);
            padding: 0.2rem 0.6rem;
            border-radius: 6px;
        }}

        .drawer-meta {{
            font-size: 0.8rem;
            color: var(--text-muted);
        }}

        .drawer-close-btn {{
            background: none;
            border: none;
            color: var(--text-muted);
            font-size: 1.25rem;
            cursor: pointer;
            padding: 0.2rem 0.5rem;
            border-radius: 6px;
            transition: all 0.2s;
            line-height: 1;
        }}

        .drawer-close-btn:hover {{
            color: var(--text-main);
            background: rgba(255, 255, 255, 0.1);
        }}

        .drawer-body {{
            flex: 1;
            overflow-y: auto;
            padding: 1.4rem;
            display: flex;
            flex-direction: column;
            gap: 1.1rem;
        }}

        .drawer-title {{
            font-size: 1.25rem;
            font-weight: 700;
            color: var(--text-main);
            line-height: 1.45;
        }}

        .drawer-connected-section, .drawer-cards-section {{
            display: flex;
            flex-direction: column;
            gap: 0.6rem;
            padding-top: 0.8rem;
            border-top: 1px dashed var(--border);
        }}

        .drawer-section-label {{
            font-size: 0.8rem;
            font-weight: 700;
            color: var(--text-muted);
            text-transform: uppercase;
            letter-spacing: 0.03em;
        }}

        .drawer-connected-chips {{
            display: flex;
            flex-wrap: wrap;
            gap: 0.45rem;
        }}

        .drawer-conn-chip {{
            font-size: 0.78rem;
            font-weight: 600;
            padding: 0.25rem 0.65rem;
            border-radius: 99px;
            border: 1px solid var(--border);
            background: rgba(99, 102, 241, 0.1);
            color: var(--primary);
            cursor: pointer;
            transition: all 0.2s;
            font-family: inherit;
        }}

        .drawer-conn-chip:hover {{
            background: var(--primary);
            color: #ffffff;
            border-color: var(--primary);
            transform: translateY(-1px);
        }}

        .drawer-cards-list {{
            display: flex;
            flex-direction: column;
            gap: 0.6rem;
            max-height: 280px;
            overflow-y: auto;
            padding-right: 0.3rem;
        }}

        .drawer-card-item {{
            background: rgba(0, 0, 0, 0.15);
            border: 1px solid var(--border);
            border-radius: 8px;
            padding: 0.65rem 0.85rem;
            cursor: pointer;
            transition: all 0.2s;
        }}

        .light-mode .drawer-card-item {{
            background: rgba(0, 0, 0, 0.02);
        }}

        .drawer-card-item:hover {{
            border-color: var(--primary);
            background: rgba(99, 102, 241, 0.08);
            transform: translateX(3px);
        }}

        .drawer-card-item-top {{
            display: flex;
            align-items: center;
            justify-content: space-between;
            margin-bottom: 0.3rem;
            font-size: 0.75rem;
        }}

        .drawer-card-key {{
            font-weight: 700;
            color: var(--primary);
        }}

        .drawer-card-jump {{
            color: var(--text-muted);
            font-size: 0.72rem;
        }}

        .drawer-card-mandarin {{
            font-size: 0.82rem;
            color: var(--text-main);
            line-height: 1.4;
            display: -webkit-box;
            -webkit-line-clamp: 2;
            -webkit-box-orient: vertical;
            overflow: hidden;
        }}

        @media (max-width: 768px) {{
            .graph-canvas-container {{
                height: 520px;
            }}
            .graph-drawer {{
                width: 100%;
                max-width: 100%;
            }}
            .graph-legend {{
                display: none;
            }}
        }}

        /* Footer block */
        footer {{
            text-align: center;
            padding: 2.5rem 1.5rem;
            border-top: 1px solid var(--border);
            color: var(--text-muted);
            font-size: 0.85rem;
            margin-top: auto;
        }}

        .update-badge {{
            display: inline-flex;
            align-items: center;
            gap: 0.35rem;
            font-size: 0.78rem;
            font-weight: 500;
            color: var(--text-muted);
            background: rgba(99, 102, 241, 0.12);
            border: 1px solid rgba(99, 102, 241, 0.25);
            padding: 0.2rem 0.65rem;
            border-radius: 99px;
            letter-spacing: 0.02em;
            vertical-align: middle;
        }}

        .light-mode .update-badge {{
            background: rgba(99, 102, 241, 0.08);
            border-color: rgba(99, 102, 241, 0.2);
            color: #4f46e5;
        }}

        /* Visibility Controls & Toggles */
        .display-toggles-bar {{
            max-width: 1200px;
            width: 100%;
            margin: 1.2rem auto 0;
            padding-top: 0.9rem;
            border-top: 1px dashed var(--border);
            display: flex;
            flex-wrap: wrap;
            align-items: center;
            justify-content: space-between;
            gap: 0.75rem;
            font-size: 0.85rem;
        }}

        .toggle-group-left {{
            display: flex;
            flex-wrap: wrap;
            align-items: center;
            gap: 0.5rem;
        }}

        .toggle-group-label {{
            color: var(--text-muted);
            font-weight: 600;
            font-size: 0.82rem;
            display: flex;
            align-items: center;
            gap: 0.3rem;
            margin-right: 0.25rem;
        }}

        .toggle-chips {{
            display: flex;
            flex-wrap: wrap;
            gap: 0.4rem;
            align-items: center;
        }}

        .toggle-chip {{
            display: inline-flex;
            align-items: center;
            gap: 0.45rem;
            padding: 0.35rem 0.8rem;
            border-radius: 99px;
            border: 1px solid var(--border);
            background: rgba(0, 0, 0, 0.2);
            color: var(--text-muted);
            font-size: 0.82rem;
            font-weight: 600;
            cursor: pointer;
            transition: all 0.2s cubic-bezier(0.4, 0, 0.2, 1);
            user-select: none;
        }}

        .light-mode .toggle-chip {{
            background: rgba(0, 0, 0, 0.04);
        }}

        .chip-dot {{
            width: 7px;
            height: 7px;
            border-radius: 50%;
            background: #64748b;
            transition: all 0.2s;
        }}

        .toggle-chip.active {{
            border-color: var(--primary);
            color: var(--text-main);
            background: rgba(99, 102, 241, 0.15);
        }}

        .light-mode .toggle-chip.active {{
            background: rgba(79, 70, 229, 0.1);
            color: var(--primary);
        }}

        .toggle-chip.active .chip-dot {{
            background: #10b981;
            box-shadow: 0 0 6px rgba(16, 185, 129, 0.6);
        }}

        .toggle-chip:hover {{
            border-color: var(--primary-hover);
            transform: translateY(-1px);
        }}

        .preset-links {{
            display: flex;
            align-items: center;
            gap: 0.4rem;
            color: var(--text-muted);
            font-size: 0.8rem;
        }}

        .preset-btn {{
            cursor: pointer;
            padding: 0.25rem 0.6rem;
            border-radius: 6px;
            background: rgba(255, 255, 255, 0.04);
            border: 1px solid var(--border);
            color: var(--text-muted);
            font-size: 0.78rem;
            font-weight: 600;
            transition: all 0.2s;
        }}

        .light-mode .preset-btn {{
            background: rgba(0, 0, 0, 0.03);
        }}

        .preset-btn:hover {{
            color: var(--primary);
            border-color: var(--primary);
            background: rgba(99, 102, 241, 0.1);
        }}

        .front-placeholder {{
            display: none;
            font-size: 1.05rem;
            color: var(--text-muted);
            font-style: italic;
            padding: 1.5rem;
            text-align: center;
            line-height: 1.6;
        }}

        /* Dynamic field hiding rules */
        body.hide-original .original-eng,
        body.hide-original #randomFrontOriginal {{
            display: none !important;
        }}

        body.hide-original .front-placeholder {{
            display: block !important;
        }}

        body.hide-mandarin .lang-mandarin,
        body.hide-mandarin #randomBackMandarinSection {{
            display: none !important;
        }}

        body.hide-english .lang-english,
        body.hide-english #randomBackEnglishSection {{
            display: none !important;
        }}

        body.hide-taiwanese .lang-taiwanese,
        body.hide-taiwanese #randomBackTaiwaneseSection {{
            display: none !important;
        }}

        body.hide-principles .card-principle-badge,
        body.hide-principles .ref-items-container,
        body.hide-principles #randomFrontPrinciples,
        body.hide-principles #randomBackPrinciples {{
            display: none !important;
        }}

        .all-hidden-notice {{
            display: none;
            padding: 1.2rem;
            margin: 0.5rem 0;
            text-align: center;
            color: var(--text-muted);
            font-size: 0.85rem;
            border: 1px dashed var(--border);
            border-radius: 8px;
            background: rgba(0, 0, 0, 0.1);
        }}

        body.hide-mandarin.hide-english.hide-taiwanese #view-principles .all-hidden-notice,
        body.hide-mandarin.hide-english.hide-taiwanese #randomCardContainer .back-body .all-hidden-notice {{
            display: block !important;
        }}

        body.hide-original.hide-mandarin.hide-english.hide-taiwanese #view-slides .all-hidden-notice {{
            display: block !important;
        }}

        /* RWD Responsive Design */
        @media (max-width: 992px) {{
            .card-grid {{
                grid-template-columns: repeat(auto-fill, minmax(300px, 1fr));
                gap: 1.2rem;
            }}
        }}

        @media (max-width: 768px) {{
            header {{
                padding: 1.25rem 1rem 1rem;
            }}
            header h1 {{
                font-size: 1.65rem;
            }}
            header p {{
                font-size: 0.88rem;
                margin-bottom: 1rem;
            }}
            .controls-container {{
                flex-direction: column;
                align-items: stretch;
                gap: 0.75rem;
            }}
            .action-tools {{
                display: flex;
                flex-wrap: wrap;
                align-items: center;
                gap: 0.6rem;
                width: 100%;
            }}
            .tabs {{
                width: 100%;
                display: flex;
                overflow-x: auto;
                white-space: nowrap;
                justify-content: space-between;
                -webkit-overflow-scrolling: touch;
            }}
            .tab-btn {{
                padding: 0.5rem 0.85rem;
                font-size: 0.85rem;
                flex: 1;
                text-align: center;
                white-space: nowrap;
            }}
            .search-bar {{
                flex: 1;
                min-width: 180px;
            }}
            .filter-tags {{
                overflow-x: auto;
                flex-wrap: nowrap;
                white-space: nowrap;
                padding-bottom: 0.4rem;
                margin-bottom: 1rem;
                -webkit-overflow-scrolling: touch;
                scrollbar-width: none;
            }}
            .filter-tags::-webkit-scrollbar {{
                display: none;
            }}
            .filter-tag {{
                flex-shrink: 0;
            }}
            .card-grid {{
                grid-template-columns: 1fr;
                gap: 1rem;
            }}
            .principle-card {{
                padding: 1.25rem 1.1rem;
            }}
            .card-title {{
                font-size: 1.1rem;
            }}
            main {{
                margin: 1.2rem auto;
                padding: 0 1rem;
            }}
            .display-toggles-bar {{
                flex-direction: column;
                align-items: flex-start;
                gap: 0.65rem;
                padding-top: 0.75rem;
            }}
            .toggle-group-left {{
                width: 100%;
            }}
            .toggle-chips {{
                width: 100%;
                overflow-x: auto;
                flex-wrap: nowrap;
                white-space: nowrap;
                padding-bottom: 0.25rem;
                -webkit-overflow-scrolling: touch;
                scrollbar-width: none;
            }}
            .toggle-chips::-webkit-scrollbar {{
                display: none;
            }}
            .toggle-chip {{
                flex-shrink: 0;
            }}
            .preset-links {{
                width: 100%;
                justify-content: flex-start;
                margin-top: 0.15rem;
            }}
            /* Random card mobile optimizations */
            .flip-card-container {{
                height: 480px;
                max-width: 100%;
            }}
            .card-face {{
                padding: 1.3rem 1.1rem;
            }}
            .front-title {{
                font-size: 1.2rem;
                line-height: 1.5;
            }}
            .back-body {{
                overflow-y: auto;
                padding-right: 0.2rem;
                gap: 0.85rem;
            }}
            .lang-section {{
                font-size: 0.9rem;
            }}
            .draw-btn {{
                width: 100%;
                justify-content: center;
                padding: 0.85rem 1.5rem;
            }}
        }}

        @media (max-width: 480px) {{
            header h1 {{
                font-size: 1.35rem;
                display: flex;
                flex-direction: column;
                gap: 0.15rem;
            }}
            .en-title {{
                font-size: 1rem;
            }}
            .header-sub-en {{
                font-size: 0.78rem;
            }}
            .principle-card {{
                padding: 1.1rem 0.9rem;
            }}
            .card-num, .card-principle-badge {{
                font-size: 0.72rem;
                padding: 0.2rem 0.45rem;
            }}
            .ref-item-chip {{
                font-size: 0.72rem;
                padding: 0.2rem 0.45rem;
            }}
            .flip-card-container {{
                height: 500px;
            }}
            .card-face {{
                padding: 1.1rem 0.85rem;
            }}
        }}
    </style>
</head>
<body>

    <header>
        <div class="controls-container">
            <div>
                <div style="display: flex; align-items: center; gap: 0.75rem; flex-wrap: wrap; margin-bottom: 0.25rem;">
                    <h1>工作的管見 <span class="en-title">My Two Cents</span></h1>
                    <span class="update-badge" title="最後建置與更新時間">🕒 最後更新：{build_timestamp}</span>
                </div>
                <p>職場生存與成長的避坑指南 (中・英・台三語對照卡牌)<span class="header-sub-en">From the Floor: My Two Cents on Practical Work Principles</span></p>
            </div>
            
            <div class="action-tools">
                <div class="tabs">
                    <button class="tab-btn active" onclick="switchView('principles')">核心原則 (20)</button>
                    <button class="tab-btn" onclick="switchView('slides')">簡報卡牌 ({len(final_items)})</button>
                    <button class="tab-btn" onclick="switchView('random')">隨機抽卡</button>
                    <button class="tab-btn" onclick="switchView('graph')">關係圖譜 🕸️</button>
                </div>
                
                <div class="search-bar" id="searchBarContainer">
                    <span class="search-icon">🔍</span>
                    <input type="text" class="search-input" id="searchInput" placeholder="搜尋原則、中文、英文或台語..." oninput="handleSearch()">
                </div>
                
                <button class="icon-btn" onclick="toggleDarkMode()" title="切換深淺色模式">🌓</button>
            </div>
        </div>
        
        <div class="display-toggles-bar">
            <div class="toggle-group-left">
                <span class="toggle-group-label">👁️ 顯示切換：</span>
                <div class="toggle-chips">
                    <button class="toggle-chip active" id="toggle-original" onclick="toggleField('original')" title="顯示/關閉投影片原文（目前書寫的原則）">
                        <span class="chip-dot"></span> 簡報原文
                    </button>
                    <button class="toggle-chip active" id="toggle-mandarin" onclick="toggleField('mandarin')" title="顯示/關閉優化中文（國語）">
                        <span class="chip-dot"></span> 優化中文
                    </button>
                    <button class="toggle-chip active" id="toggle-english" onclick="toggleField('english')" title="顯示/關閉優化英文">
                        <span class="chip-dot"></span> 優化英文
                    </button>
                    <button class="toggle-chip active" id="toggle-taiwanese" onclick="toggleField('taiwanese')" title="顯示/關閉優化台文">
                        <span class="chip-dot"></span> 優化台文
                    </button>
                    <button class="toggle-chip active" id="toggle-principles" onclick="toggleField('principles')" title="顯示/關閉核心原則標籤與對照連結">
                        <span class="chip-dot"></span> 對照原則
                    </button>
                </div>
            </div>
            <div class="preset-links">
                <span>快速模式：</span>
                <button class="preset-btn" onclick="applyPreset('all')" title="顯示全部內容">全部顯示</button>
                <button class="preset-btn" onclick="applyPreset('trilingual')" title="只看三語翻譯，隱藏簡報原文">純三語</button>
                <button class="preset-btn" onclick="applyPreset('original_only')" title="只看簡報原文，隱藏三語翻譯">純原文</button>
            </div>
        </div>
    </header>

    <main>
        <!-- VIEW 1: Core Principles (20) -->
        <section id="view-principles" class="section-view active">
            <div class="filter-tags" id="principlesFilters">
                <button class="filter-tag active" onclick="filterPrinciples('all')">全部章節</button>
                <button class="filter-tag" onclick="filterPrinciples('1')">第一篇：效能與工具</button>
                <button class="filter-tag" onclick="filterPrinciples('2')">第二篇：思考與決策</button>
                <button class="filter-tag" onclick="filterPrinciples('3')">第三篇：溝通與協作</button>
                <button class="filter-tag" onclick="filterPrinciples('4')">第四篇：職場生存與防線</button>
                <button class="filter-tag" onclick="filterPrinciples('5')">第五篇：職涯與個人成長</button>
                <button class="filter-tag" onclick="filterPrinciples('6')">第六篇：能量、韌性與生活平衡</button>
            </div>
            <div class="card-grid" id="principlesGrid"></div>
        </section>

        <!-- VIEW 2: Slide Items -->
        <section id="view-slides" class="section-view">
            <div class="filter-tags" id="slidesFilters">
                {slides_filter_html}
            </div>
            <div class="card-grid" id="slidesGrid"></div>
        </section>

        <!-- VIEW 3: Random Card -->
        <section id="view-random" class="section-view">
            <div class="random-view-container">
                <div class="flip-card-container" id="randomCardContainer" onclick="flipRandomCard()">
                    <div class="flip-card">
                        <!-- Front Face -->
                        <div class="card-face card-face-front">
                            <div class="card-header">
                                <div style="display: flex; align-items: center; flex-wrap: wrap; gap: 0.25rem;">
                                    <span class="card-num" id="randomFrontNum">No. 1</span>
                                    <span id="randomFrontPrinciples"></span>
                                </div>
                                <span class="card-meta" id="randomFrontMeta">Slide 1</span>
                            </div>
                            <div class="front-body">
                                <p class="front-title" id="randomFrontOriginal">Click 'Draw Card' to start!</p>
                                <p class="front-placeholder">（簡報原文目前已隱藏，點擊卡牌翻面查看三語對照）</p>
                            </div>
                            <p class="flip-hint">💡 點擊卡牌翻面查看三語對照</p>
                        </div>
                        <!-- Back Face -->
                        <div class="card-face card-face-back">
                            <div class="card-header">
                                <div style="display: flex; align-items: center; flex-wrap: wrap; gap: 0.25rem;">
                                    <span class="card-num" id="randomBackNum">No. 1</span>
                                    <span id="randomBackPrinciples"></span>
                                </div>
                                <span class="card-meta" id="randomBackMeta">三語對照</span>
                            </div>
                            <div class="back-body">
                                <div class="lang-section lang-mandarin" id="randomBackMandarinSection">
                                    <span class="lang-label">優化中文 (Polished Mandarin)</span>
                                    <p class="lang-txt" id="randomBackMandarin">中文翻譯與潤稿內容</p>
                                </div>
                                <div class="lang-section lang-english" id="randomBackEnglishSection">
                                    <span class="lang-label">優化英文 (Polished English)</span>
                                    <p class="lang-txt" id="randomBackEnglish">English polished text</p>
                                </div>
                                <div class="lang-section lang-taiwanese" id="randomBackTaiwaneseSection">
                                    <span class="lang-label">優化台文 (Taiwanese)</span>
                                    <p class="lang-txt" id="randomBackTaiwanese">台語漢字智慧</p>
                                </div>
                                <div class="all-hidden-notice">⚠️ 所有三語翻譯欄位均已關閉，請由上方「顯示切換」開啟</div>
                            </div>
                            <p class="flip-hint">💡 點擊卡牌翻回正面</p>
                        </div>
                    </div>
                </div>
                <button class="draw-btn" onclick="drawRandomCard()">
                    <span>🃏</span> 隨機抽一張卡
                </button>
            </div>
        </section>

        <!-- VIEW 4: Connection Graph (關係圖譜) -->
        <section id="view-graph" class="section-view">
            <div class="graph-wrapper">
                <div class="graph-toolbar">
                    <div class="graph-toolbar-left">
                        <span class="graph-title">🕸️ 知識關聯圖譜</span>
                        <div class="graph-mode-switch" id="graphModeSwitch">
                            <button class="graph-mode-btn active" id="btnModePrinciples" onclick="setGraphMode('principles')">🌐 原則骨幹 (20)</button>
                            <button class="graph-mode-btn" id="btnModeGalaxy" onclick="setGraphMode('galaxy')">✨ 全景星系 (458卡牌)</button>
                        </div>
                        <div class="graph-filters" id="graphFilterChips">
                            <button class="filter-tag active" onclick="filterGraphChapter('all')">全部章節</button>
                            <button class="filter-tag" onclick="filterGraphChapter('1')">第一篇</button>
                            <button class="filter-tag" onclick="filterGraphChapter('2')">第二篇</button>
                            <button class="filter-tag" onclick="filterGraphChapter('3')">第三篇</button>
                            <button class="filter-tag" onclick="filterGraphChapter('4')">第四篇</button>
                            <button class="filter-tag" onclick="filterGraphChapter('5')">第五篇</button>
                            <button class="filter-tag" onclick="filterGraphChapter('6')">第六篇</button>
                        </div>
                    </div>
                    <div class="graph-toolbar-right">
                        <button class="graph-btn" id="btnGraphFreeze" onclick="toggleGraphSimulation()" title="暫停 / 恢復節點物理模擬">⏸️ 暫停模擬</button>
                        <button class="graph-btn" onclick="resetGraphView()" title="重置視角與縮放">🎯 重置視角</button>
                    </div>
                </div>
                
                <div class="graph-canvas-container" id="graphCanvasContainer">
                    <canvas id="graphCanvas"></canvas>
                    
                    <div class="graph-legend">
                        <div class="legend-chip"><span class="legend-dot" style="background:#6366f1;"></span>Ch.1 效能工具</div>
                        <div class="legend-chip"><span class="legend-dot" style="background:#0ea5e9;"></span>Ch.2 思考決策</div>
                        <div class="legend-chip"><span class="legend-dot" style="background:#10b981;"></span>Ch.3 溝通協作</div>
                        <div class="legend-chip"><span class="legend-dot" style="background:#f59e0b;"></span>Ch.4 職場生存</div>
                        <div class="legend-chip"><span class="legend-dot" style="background:#8b5cf6;"></span>Ch.5 職涯成長</div>
                        <div class="legend-chip"><span class="legend-dot" style="background:#ec4899;"></span>Ch.6 韌性平衡</div>
                        <div class="legend-chip"><span class="legend-dot" style="background:#fbbf24; border: 1px solid #ffffff;"></span>跨界交匯卡</div>
                    </div>

                    <div class="graph-hint">
                        💡 拖曳平移 · 滾輪縮放 · 雙擊原則展開/收合卡牌 · 點擊原則或卡牌查看詳情
                    </div>

                    <!-- Slide-in Detail Drawer -->
                    <div class="graph-drawer" id="graphDrawer">
                        <div class="drawer-header">
                            <div class="drawer-title-group">
                                <span class="drawer-badge" id="drawerBadge">原則 1</span>
                                <span class="drawer-meta" id="drawerMeta">第一篇：效能與工具</span>
                            </div>
                            <div style="display: flex; align-items: center; gap: 0.4rem;">
                                <button class="drawer-action-btn" id="btnDrawerSatellite" onclick="toggleCurrentPrincipleSatellite()" title="在畫布上展開或收合此原則的卡牌衛星">🪐 展開卡牌</button>
                                <button class="drawer-close-btn" onclick="closeGraphDrawer()" title="關閉面板">✕</button>
                            </div>
                        </div>
                        <div class="drawer-body">
                            <h3 class="drawer-title" id="drawerTitle">原則標題</h3>
                            <div class="lang-section lang-mandarin">
                                <span class="lang-label">國語 (Mandarin)</span>
                                <p class="lang-txt" id="drawerMandarin"></p>
                            </div>
                            <div class="lang-section lang-english">
                                <span class="lang-label">English</span>
                                <p class="lang-txt" id="drawerEnglish"></p>
                            </div>
                            <div class="lang-section lang-taiwanese">
                                <span class="lang-label">台語 (Taiwanese)</span>
                                <p class="lang-txt" id="drawerTaiwanese"></p>
                            </div>

                            <div class="drawer-connected-section">
                                <span class="drawer-section-label">🔗 關聯原則 (點擊聚焦)</span>
                                <div class="drawer-connected-chips" id="drawerConnectedPrinciples"></div>
                            </div>

                            <div class="drawer-cards-section">
                                <span class="drawer-section-label" id="drawerCardsLabel">📑 對照簡報卡牌 (0 則)</span>
                                <div class="drawer-cards-list" id="drawerCardsList"></div>
                            </div>
                        </div>
                    </div>
                </div>
            </div>
        </section>
    </main>

    <footer>
        <p>© 2026 工作的管見 - 避坑指南與商務工作原則整理與潤稿系統 · 最後更新：{build_timestamp}</p>
    </footer>

    <script>
        // Embed parsed JSON data
        const PRINCIPLES = {principles_json};
        const SLIDES = {slides_json};
        const CHAPTERS = {chapters_json};

        let currentActiveView = 'principles';
        let principlesFilter = 'all';
        let slidesFilter = 'all';
        let searchQuery = '';

        const DEFAULT_DISPLAY = {{
            original: true,
            principles: true,
            mandarin: true,
            english: true,
            taiwanese: true
        }};

        let displaySettings = {{ ...DEFAULT_DISPLAY }};

        function loadDisplaySettings() {{
            try {{
                const saved = localStorage.getItem('work_principles_display');
                if (saved) {{
                    displaySettings = Object.assign({{}}, DEFAULT_DISPLAY, JSON.parse(saved));
                }}
            }} catch(e) {{}}
            applyDisplayClasses();
        }}

        function saveDisplaySettings() {{
            try {{
                localStorage.setItem('work_principles_display', JSON.stringify(displaySettings));
            }} catch(e) {{}}
        }}

        function toggleField(field) {{
            displaySettings[field] = !displaySettings[field];
            saveDisplaySettings();
            applyDisplayClasses();
        }}

        function applyPreset(type) {{
            if (type === 'all') {{
                displaySettings = {{ original: true, principles: true, mandarin: true, english: true, taiwanese: true }};
            }} else if (type === 'trilingual') {{
                displaySettings = {{ original: false, principles: true, mandarin: true, english: true, taiwanese: true }};
            }} else if (type === 'original_only') {{
                displaySettings = {{ original: true, principles: true, mandarin: false, english: false, taiwanese: false }};
            }}
            saveDisplaySettings();
            applyDisplayClasses();
        }}

        function applyDisplayClasses() {{
            const fields = ['original', 'principles', 'mandarin', 'english', 'taiwanese'];
            fields.forEach(f => {{
                const btn = document.getElementById(`toggle-${{f}}`);
                if (displaySettings[f]) {{
                    document.body.classList.remove(`hide-${{f}}`);
                    if (btn) btn.classList.add('active');
                }} else {{
                    document.body.classList.add(`hide-${{f}}`);
                    if (btn) btn.classList.remove('active');
                }}
            }});
        }}

        // Initialize UI
        document.addEventListener('DOMContentLoaded', () => {{
            loadDisplaySettings();
            renderPrinciples();
            renderSlides();
            drawRandomCard();
        }});

        function switchView(viewName) {{
            currentActiveView = viewName;
            
            // Toggle active tabs
            const tabButtons = document.querySelectorAll('.tab-btn');
            tabButtons.forEach(btn => btn.classList.remove('active'));
            
            const tabIndexMap = {{ 'principles': 0, 'slides': 1, 'random': 2, 'graph': 3 }};
            if (tabIndexMap[viewName] !== undefined && tabButtons[tabIndexMap[viewName]]) {{
                tabButtons[tabIndexMap[viewName]].classList.add('active');
            }}
            
            // Toggle view containers
            document.querySelectorAll('.section-view').forEach(view => view.classList.remove('active'));
            const targetView = document.getElementById(`view-${{viewName}}`);
            if (targetView) targetView.classList.add('active');
            
            // Show/hide search bar based on view
            const searchBar = document.getElementById('searchBarContainer');
            if (viewName === 'random' || viewName === 'graph') {{
                searchBar.style.display = 'none';
            }} else {{
                searchBar.style.display = 'flex';
                // Reset search query
                document.getElementById('searchInput').value = '';
                searchQuery = '';
                handleSearch();
            }}

            if (viewName === 'graph') {{
                initOrResizeGraph();
            }} else {{
                if (typeof graphAnimationId !== 'undefined' && graphAnimationId) {{
                    cancelAnimationFrame(graphAnimationId);
                    graphAnimationId = null;
                }}
            }}
        }}

        function toggleDarkMode() {{
            document.body.classList.toggle('light-mode');
        }}

        function getChapterName(num) {{
            if (num <= 4) return '第一篇：效能與工具';
            if (num <= 8) return '第二篇：思考與決策';
            if (num <= 11) return '第三篇：溝通與協作';
            if (num <= 14) return '第四篇：職場生存與防線';
            if (num <= 17) return '第五篇：職涯與個人成長';
            return '第六篇：能量、韌性與生活平衡';
        }}

        function getChapterId(num) {{
            if (num <= 4) return '1';
            if (num <= 8) return '2';
            if (num <= 11) return '3';
            if (num <= 14) return '4';
            if (num <= 17) return '5';
            return '6';
        }}

        function getMatchedPrinciples(slideKey) {{
            const baseKey = slideKey.split('-')[0];
            const matched = [];
            PRINCIPLES.forEach(p => {{
                const refs = p.ref.split(',').map(r => r.trim());
                if (refs.includes(slideKey) || refs.includes(baseKey)) {{
                    matched.push(p.num);
                }}
            }});
            return matched;
        }}

        function renderPrinciples() {{
            const grid = document.getElementById('principlesGrid');
            grid.innerHTML = '';
            
            const filtered = PRINCIPLES.filter(p => {{
                const chapterId = getChapterId(p.num);
                const matchesChapter = (principlesFilter === 'all' || principlesFilter === chapterId);
                const matchesSearch = !searchQuery || 
                    p.title.toLowerCase().includes(searchQuery) ||
                    p.mandarin.toLowerCase().includes(searchQuery) ||
                    p.english.toLowerCase().includes(searchQuery) ||
                    p.taiwanese.toLowerCase().includes(searchQuery);
                return matchesChapter && matchesSearch;
            }});

            if (filtered.length === 0) {{
                grid.innerHTML = `<div style="grid-column: 1/-1; text-align: center; color: var(--text-muted); padding: 3rem;">無符合搜尋條件的原則</div>`;
                return;
            }}

            // Render by Chapter Groups with Chapter Banner and Golden Quotes
            CHAPTERS.forEach(chap => {{
                const chapPrinciples = filtered.filter(p => getChapterId(p.num) === chap.id);
                if (chapPrinciples.length === 0) return;

                // 1. Chapter Header Banner with Golden Quotes
                const banner = document.createElement('div');
                banner.className = 'chapter-header-banner';

                const quotesPillsHTML = chap.quotes.map(q => 
                    `<span class="quote-pill" title="篇章核心指引金句">${{q}}</span>`
                ).join(' ');

                const minNum = Math.min(...chapPrinciples.map(p => p.num));
                const maxNum = Math.max(...chapPrinciples.map(p => p.num));
                const rangeText = minNum === maxNum ? `原則 ${{minNum}}` : `原則 ${{minNum}} - ${{maxNum}}`;

                banner.innerHTML = `
                    <div class="chapter-banner-top">
                        <span class="chapter-badge">Chapter 0${{chap.id}}</span>
                        <h2 class="chapter-title">${{chap.title}}</h2>
                        <span class="chapter-range-chip">${{rangeText}} (共 ${{chapPrinciples.length}} 條原則)</span>
                    </div>
                    <div class="chapter-quotes-area">
                        <div class="chapter-quotes-header">
                            <span>✨ 篇章核心指引金句 (${{chap.quotes.length}} 則)</span>
                        </div>
                        <div class="chapter-quotes-pills">
                            ${{quotesPillsHTML}}
                        </div>
                    </div>
                `;
                grid.appendChild(banner);

                // 2. Principle Cards for this Chapter
                chapPrinciples.forEach(p => {{
                    const card = document.createElement('div');
                    card.className = 'principle-card';
                    card.id = `card-principle-${{p.num}}`;

                    const refKeys = p.ref.split(',').map(r => r.trim()).filter(Boolean);
                    const refChipsHTML = refKeys.map(key => 
                        `<button class="ref-item-chip" onclick="jumpToSlideCard('${{key}}')" title="點擊直接查看簡報卡牌 No. ${{key}}">No. ${{key}}</button>`
                    ).join(' ');

                    card.innerHTML = `
                        <div class="card-header">
                            <span class="card-num">原則 ${{p.num}}</span>
                            <span class="card-meta">${{getChapterName(p.num)}}</span>
                        </div>
                        <h3 class="card-title">${{p.title}}</h3>
                        <div class="lang-section lang-mandarin">
                            <span class="lang-label">國語 (Mandarin)</span>
                            <p class="lang-txt">${{p.mandarin}}</p>
                        </div>
                        <div class="lang-section lang-english">
                            <span class="lang-label">English</span>
                            <p class="lang-txt">${{p.english}}</p>
                        </div>
                        <div class="lang-section lang-taiwanese">
                            <span class="lang-label">台語 (Taiwanese)</span>
                            <p class="lang-txt">${{p.taiwanese}}</p>
                        </div>
                        <div class="all-hidden-notice">⚠️ 所有三語翻譯欄位均已關閉</div>
                        <div class="ref-items-container">
                            <span class="ref-items-label">🔗 對照筆記卡牌 (${{refKeys.length}} 條項目，點擊跳轉)</span>
                            <div class="ref-chips-grid">
                                ${{refChipsHTML}}
                            </div>
                        </div>
                    `;
                    grid.appendChild(card);
                }});
            }});
        }}

        function renderSlides() {{
            const grid = document.getElementById('slidesGrid');
            grid.innerHTML = '';

            const filtered = SLIDES.filter(s => {{
                const matchesSlide = (slidesFilter === 'all' || slidesFilter === String(s.slide));
                
                // Get matched principles for this slide
                const matched = getMatchedPrinciples(s.key);
                const matchesPrincipleSearch = matched.some(num => {{
                    const p = PRINCIPLES.find(pr => pr.num === num);
                    if (!p) return false;
                    const searchLower = searchQuery.toLowerCase();
                    return `原則 ${{num}}`.toLowerCase().includes(searchLower) ||
                           `原則${{num}}`.toLowerCase().includes(searchLower) ||
                           p.title.toLowerCase().includes(searchLower);
                }});

                const matchesSearch = !searchQuery || 
                    s.key.toLowerCase().includes(searchQuery) ||
                    s.original.toLowerCase().includes(searchQuery) ||
                    s.mandarin.toLowerCase().includes(searchQuery) ||
                    s.english.toLowerCase().includes(searchQuery) ||
                    s.taiwanese.toLowerCase().includes(searchQuery) ||
                    matchesPrincipleSearch;
                return matchesSlide && matchesSearch;
            }});

            if (filtered.length === 0) {{
                grid.innerHTML = `<div style="grid-column: 1/-1; text-align: center; color: var(--text-muted); padding: 3rem;">無符合搜尋條件的卡牌</div>`;
                return;
            }}

            filtered.forEach(s => {{
                const card = document.createElement('div');
                card.className = 'principle-card';
                card.id = `card-slide-${{s.key}}`;
                const matched = getMatchedPrinciples(s.key);
                const badges = matched.map(num => `<span class="card-principle-badge" onclick="jumpToPrinciple(${{num}})" style="cursor:pointer;" title="點擊跳轉至 原則 ${{num}}">原則 ${{num}}</span>`).join(' ');
                card.innerHTML = `
                    <div class="card-header">
                        <div style="display: flex; align-items: center; flex-wrap: wrap; gap: 0.25rem;">
                            <span class="card-num">No. ${{s.key}}</span>
                            ${{badges}}
                        </div>
                        <span class="card-meta">Slide ${{s.slide}}</span>
                    </div>
                    <div class="original-eng">"${{s.original}}"</div>
                    <div class="lang-section lang-mandarin">
                        <span class="lang-label">優化中文 (Polished Mandarin)</span>
                        <p class="lang-txt">${{s.mandarin}}</p>
                    </div>
                    <div class="lang-section lang-english">
                        <span class="lang-label">優化英文 (Polished English)</span>
                        <p class="lang-txt">${{s.english}}</p>
                    </div>
                    <div class="lang-section lang-taiwanese">
                        <span class="lang-label">優化台文 (Taiwanese)</span>
                        <p class="lang-txt">${{s.taiwanese}}</p>
                    </div>
                    <div class="all-hidden-notice">⚠️ 簡報原文與三語翻譯均已關閉</div>
                `;
                grid.appendChild(card);
            }});
        }}

        function jumpToSlideCard(key) {{
            switchView('slides');
            
            slidesFilter = 'all';
            document.querySelectorAll('#slidesFilters .filter-tag').forEach(tag => tag.classList.remove('active'));
            const allTag = document.querySelector('#slidesFilters .filter-tag');
            if (allTag) allTag.classList.add('active');
            
            searchQuery = '';
            const searchInput = document.getElementById('searchInput');
            if (searchInput) searchInput.value = '';
            
            renderSlides();
            
            setTimeout(() => {{
                const targetCard = document.getElementById(`card-slide-${{key}}`);
                if (targetCard) {{
                    targetCard.scrollIntoView({{ behavior: 'smooth', block: 'center' }});
                    targetCard.classList.add('highlight-pulse');
                    setTimeout(() => {{
                        targetCard.classList.remove('highlight-pulse');
                    }}, 2500);
                }}
            }}, 120);
        }}

        function jumpToPrinciple(num) {{
            switchView('principles');
            principlesFilter = 'all';
            document.querySelectorAll('#principlesFilters .filter-tag').forEach(tag => tag.classList.remove('active'));
            const allTag = document.querySelector('#principlesFilters .filter-tag');
            if (allTag) allTag.classList.add('active');
            
            searchQuery = '';
            const searchInput = document.getElementById('searchInput');
            if (searchInput) searchInput.value = '';
            
            renderPrinciples();

            setTimeout(() => {{
                const targetCard = document.getElementById(`card-principle-${{num}}`);
                if (targetCard) {{
                    targetCard.scrollIntoView({{ behavior: 'smooth', block: 'center' }});
                    targetCard.classList.add('highlight-pulse');
                    setTimeout(() => {{
                        targetCard.classList.remove('highlight-pulse');
                    }}, 2500);
                }}
            }}, 120);
        }}

        function filterPrinciples(chapter) {{
            principlesFilter = chapter;
            document.querySelectorAll('#principlesFilters .filter-tag').forEach(tag => tag.classList.remove('active'));
            event.target.classList.add('active');
            renderPrinciples();
        }}

        function filterSlides(slide) {{
            slidesFilter = slide;
            document.querySelectorAll('#slidesFilters .filter-tag').forEach(tag => tag.classList.remove('active'));
            event.target.classList.add('active');
            renderSlides();
        }}

        function handleSearch() {{
            searchQuery = document.getElementById('searchInput').value.toLowerCase().trim();
            if (currentActiveView === 'principles') {{
                renderPrinciples();
            }} else if (currentActiveView === 'slides') {{
                renderSlides();
            }}
        }}

        // Random Card logic
        function drawRandomCard() {{
            const container = document.getElementById('randomCardContainer');
            container.classList.remove('flipped'); // Reset flip state to show front
            
            // Choose a random slide card
            const randomIndex = Math.floor(Math.random() * SLIDES.length);
            const cardData = SLIDES[randomIndex];
            
            // Create sparkle particles
            createSparkles();

            // Animate card deal
            const innerCard = container.querySelector('.flip-card');
            innerCard.style.transform = 'scale(0.9) rotateY(-10deg)';
            
            setTimeout(() => {{
                // Update text content
                document.getElementById('randomFrontNum').innerText = `No. ${{cardData.key}}`;
                document.getElementById('randomFrontMeta').innerText = `Slide ${{cardData.slide}}`;
                document.getElementById('randomFrontOriginal').innerText = cardData.original;
                
                document.getElementById('randomBackNum').innerText = `No. ${{cardData.key}}`;
                document.getElementById('randomBackMeta').innerText = `Slide ${{cardData.slide}} ・ 對照翻譯`;
                document.getElementById('randomBackMandarin').innerText = cardData.mandarin;
                document.getElementById('randomBackEnglish').innerText = cardData.english;
                document.getElementById('randomBackTaiwanese').innerText = cardData.taiwanese;
                
                // Update matched principles badges
                const matched = getMatchedPrinciples(cardData.key);
                const badges = matched.map(num => `<span class="card-principle-badge">原則 ${{num}}</span>`).join(' ');
                document.getElementById('randomFrontPrinciples').innerHTML = badges;
                document.getElementById('randomBackPrinciples').innerHTML = badges;
                
                innerCard.style.transform = '';
            }}, 200);
        }}

        function flipRandomCard() {{
            const container = document.getElementById('randomCardContainer');
            container.classList.toggle('flipped');
        }}

        function createSparkles() {{
            const container = document.getElementById('view-random');
            const button = document.querySelector('.draw-btn');
            const rect = button.getBoundingClientRect();
            
            // Generate some particle spans
            for (let i = 0; i < 20; i++) {{
                const particle = document.createElement('div');
                particle.className = 'sparkle-particles';
                
                // Color variations
                const colors = ['#6366f1', '#818cf8', '#a5b4fc', '#8b5cf6', '#a78bfa'];
                particle.style.backgroundColor = colors[Math.floor(Math.random() * colors.length)];
                
                // Position above button
                particle.style.left = `${{rect.left + rect.width/2}}px`;
                particle.style.top = `${{rect.top + window.scrollY}}px`;
                particle.style.opacity = '1';
                particle.style.width = `${{Math.random() * 8 + 4}}px`;
                particle.style.height = particle.style.width;
                
                document.body.appendChild(particle);
                
                // Animate outward
                const angle = Math.random() * Math.PI * 2;
                const distance = Math.random() * 120 + 30;
                const destX = Math.cos(angle) * distance;
                const destY = Math.sin(angle) * distance - 80;
                
                particle.animate([
                    {{ transform: 'translate(0, 0) scale(1)', opacity: 1 }},
                    {{ transform: `translate(${{destX}}px, ${{destY}}px) scale(0)`, opacity: 0 }}
                ], {{
                    duration: Math.random() * 600 + 400,
                    easing: 'cubic-bezier(0.1, 0.8, 0.3, 1)',
                    fill: 'forwards'
                }});
                
                setTimeout(() => {{
                    particle.remove();
                }}, 1000);
            }}
        }}

        // ==========================================
        // VIEW 4: Interactive Connection Graph (關係圖譜) Engine
        // ==========================================
        const CHAPTER_THEMES = {{
            '1': {{ color: '#6366f1', name: '第一篇：效能與工具' }},
            '2': {{ color: '#0ea5e9', name: '第二篇：思考與決策' }},
            '3': {{ color: '#10b981', name: '第三篇：溝通與協作' }},
            '4': {{ color: '#f59e0b', name: '第四篇：職場生存與防線' }},
            '5': {{ color: '#8b5cf6', name: '第五篇：職涯與個人成長' }},
            '6': {{ color: '#ec4899', name: '第六篇：能量、韌性與生活平衡' }}
        }};

        const SEMANTIC_SYNERGIES = [
            [1, 19], [19, 20], [1, 20],
            [11, 13], [12, 13],
            [3, 4], [7, 8],
            [6, 18],
            [15, 17],
            [9, 14],
            [2, 5],
            [10, 16]
        ];

        let graphNodes = [];
        let graphEdges = [];
        let graphInitialized = false;
        let graphAnimationId = null;
        let isGraphSimulating = true;
        let graphChapterFilter = 'all';
        let graphMode = 'principles'; // 'principles' or 'galaxy'
        let expandedPrinciples = new Set(); // Set of principle numbers (e.g. 1, 6)

        let graphCam = {{ x: 0, y: 0, scale: 1.0 }};
        let isPanning = false;
        let panStartX = 0;
        let panStartY = 0;
        let mouseMovedDuringClick = false;
        let draggedNode = null;
        let hoveredNode = null;
        let selectedNode = null;

        function buildGraphData() {{
            const prevPos = new Map();
            graphNodes.forEach(n => prevPos.set(n.id, {{ x: n.x, y: n.y, vx: n.vx, vy: n.vy }}));

            graphNodes = [];
            graphEdges = [];

            // 1. Chapter Hub Nodes (6)
            CHAPTERS.forEach((chap, idx) => {{
                const id = chap.id;
                const angle = (idx / 6) * Math.PI * 2 - Math.PI / 2;
                const ringRadius = 280;
                let nx = Math.cos(angle) * ringRadius;
                let ny = Math.sin(angle) * ringRadius;
                let nvx = 0, nvy = 0;
                if (prevPos.has(`c${{id}}`)) {{
                    const p = prevPos.get(`c${{id}}`);
                    nx = p.x; ny = p.y; nvx = p.vx; nvy = p.vy;
                }}
                graphNodes.push({{
                    id: `c${{id}}`,
                    chapterId: id,
                    type: 'chapter',
                    num: parseInt(id),
                    title: chap.title,
                    color: CHAPTER_THEMES[id].color,
                    radius: 34,
                    mass: 5.0,
                    x: nx,
                    y: ny,
                    vx: nvx,
                    vy: nvy
                }});
            }});

            // 2. Principle Nodes (20)
            PRINCIPLES.forEach(p => {{
                const chapId = getChapterId(p.num);
                const chapNode = graphNodes.find(n => n.id === `c${{chapId}}`);
                const refCount = p.ref.split(',').map(r => r.trim()).filter(Boolean).length;
                
                const jitterAngle = Math.random() * Math.PI * 2;
                const jitterDist = 80 + Math.random() * 70;

                let nx = chapNode ? chapNode.x + Math.cos(jitterAngle) * jitterDist : (Math.random() - 0.5) * 350;
                let ny = chapNode ? chapNode.y + Math.sin(jitterAngle) * jitterDist : (Math.random() - 0.5) * 350;
                let nvx = 0, nvy = 0;
                if (prevPos.has(`p${{p.num}}`)) {{
                    const prev = prevPos.get(`p${{p.num}}`);
                    nx = prev.x; ny = prev.y; nvx = prev.vx; nvy = prev.vy;
                }}

                const pNode = {{
                    id: `p${{p.num}}`,
                    chapterId: chapId,
                    type: 'principle',
                    num: p.num,
                    title: p.title,
                    data: p,
                    color: CHAPTER_THEMES[chapId].color,
                    refCount: refCount,
                    radius: Math.min(28, Math.max(20, 18 + refCount * 0.35)),
                    mass: 1.5,
                    x: nx,
                    y: ny,
                    vx: nvx,
                    vy: nvy
                }};
                graphNodes.push(pNode);

                // Chapter -> Principle structural edge
                graphEdges.push({{
                    source: `c${{chapId}}`,
                    target: `p${{p.num}}`,
                    type: 'chapter',
                    distance: 140,
                    strength: 0.04,
                    color: CHAPTER_THEMES[chapId].color
                }});
            }});

            // 3. Shared Slide Cards Edges (Principle <-> Principle)
            for (let i = 0; i < PRINCIPLES.length; i++) {{
                const p1 = PRINCIPLES[i];
                const refs1 = new Set(p1.ref.split(',').map(r => r.trim()).filter(Boolean));
                for (let j = i + 1; j < PRINCIPLES.length; j++) {{
                    const p2 = PRINCIPLES[j];
                    const refs2 = p2.ref.split(',').map(r => r.trim()).filter(Boolean);
                    const common = refs2.filter(r => refs1.has(r));
                    if (common.length > 0) {{
                        graphEdges.push({{
                            source: `p${{p1.num}}`,
                            target: `p${{p2.num}}`,
                            type: 'shared',
                            sharedCards: common,
                            sharedCount: common.length,
                            distance: 170,
                            strength: 0.022 * Math.min(3, common.length),
                            color: '#94a3b8'
                        }});
                    }}
                }}
            }}

            // 4. Semantic Synergies (Principle <-> Principle)
            SEMANTIC_SYNERGIES.forEach(([n1, n2]) => {{
                const existing = graphEdges.find(e => 
                    (e.source === `p${{n1}}` && e.target === `p${{n2}}`) ||
                    (e.source === `p${{n2}}` && e.target === `p${{n1}}`)
                );
                if (!existing) {{
                    graphEdges.push({{
                        source: `p${{n1}}`,
                        target: `p${{n2}}`,
                        type: 'synergy',
                        distance: 220,
                        strength: 0.015,
                        color: '#c084fc'
                    }});
                }}
            }});

            // 5. Card Satellite Nodes (458 Cards)
            const activeCardKeys = new Set();
            if (graphMode === 'galaxy') {{
                SLIDES.forEach(s => activeCardKeys.add(s.key));
            }} else {{
                PRINCIPLES.forEach(p => {{
                    if (expandedPrinciples.has(p.num)) {{
                        SLIDES.forEach(s => {{
                            if (getMatchedPrinciples(s.key).includes(p.num)) {{
                                activeCardKeys.add(s.key);
                            }}
                        }});
                    }}
                }});
            }}

            activeCardKeys.forEach(key => {{
                const s = SLIDES.find(item => item.key === key);
                if (!s) return;
                const matched = getMatchedPrinciples(key);
                const primaryPNum = matched[0] || 1;
                const primaryChapId = getChapterId(primaryPNum);
                const isBridge = matched.length > 1;
                const cardColor = isBridge ? '#fbbf24' : CHAPTER_THEMES[primaryChapId].color;

                let nx, ny, nvx = 0, nvy = 0;
                if (prevPos.has(`s${{key}}`)) {{
                    const prev = prevPos.get(`s${{key}}`);
                    nx = prev.x; ny = prev.y; nvx = prev.vx; nvy = prev.vy;
                }} else {{
                    const pNode = graphNodes.find(n => n.id === `p${{primaryPNum}}`);
                    const angle = Math.random() * Math.PI * 2;
                    const dist = 55 + Math.random() * 45;
                    nx = pNode ? pNode.x + Math.cos(angle) * dist : (Math.random() - 0.5) * 400;
                    ny = pNode ? pNode.y + Math.sin(angle) * dist : (Math.random() - 0.5) * 400;
                }}

                graphNodes.push({{
                    id: `s${{key}}`,
                    type: 'card',
                    key: key,
                    num: s.num,
                    slide: s.slide,
                    title: `No. ${{key}}`,
                    data: s,
                    chapterId: primaryChapId,
                    principles: matched,
                    isBridge: isBridge,
                    color: cardColor,
                    radius: isBridge ? 8 : 6.5,
                    mass: 0.45,
                    x: nx,
                    y: ny,
                    vx: nvx,
                    vy: nvy
                }});

                // Edge from each matching principle to this card
                matched.forEach(pNum => {{
                    if (graphNodes.some(n => n.id === `p${{pNum}}`)) {{
                        graphEdges.push({{
                            source: `p${{pNum}}`,
                            target: `s${{key}}`,
                            type: 'card',
                            distance: 72,
                            strength: 0.048,
                            color: CHAPTER_THEMES[getChapterId(pNum)].color
                        }});
                    }}
                }});
            }});
        }}

        function initOrResizeGraph() {{
            const canvas = document.getElementById('graphCanvas');
            const container = document.getElementById('graphCanvasContainer');
            if (!canvas || !container) return;

            const dpr = window.devicePixelRatio || 1;
            const w = container.clientWidth;
            const h = container.clientHeight;

            canvas.width = w * dpr;
            canvas.height = h * dpr;
            canvas.style.width = `${{w}}px`;
            canvas.style.height = `${{h}}px`;

            if (!graphInitialized) {{
                buildGraphData();
                setupGraphInteractions();
                graphInitialized = true;
            }}

            if (!graphAnimationId) {{
                graphAnimationLoop();
            }}
        }}

        function getCanvasMousePos(evt) {{
            const canvas = document.getElementById('graphCanvas');
            const rect = canvas.getBoundingClientRect();
            let clientX = evt.clientX;
            let clientY = evt.clientY;
            if (evt.touches && evt.touches[0]) {{
                clientX = evt.touches[0].clientX;
                clientY = evt.touches[0].clientY;
            }} else if (evt.changedTouches && evt.changedTouches[0]) {{
                clientX = evt.changedTouches[0].clientX;
                clientY = evt.changedTouches[0].clientY;
            }}
            const sx = clientX - rect.left;
            const sy = clientY - rect.top;
            const wx = (sx - rect.width / 2 - graphCam.x) / graphCam.scale;
            const wy = (sy - rect.height / 2 - graphCam.y) / graphCam.scale;
            return {{ screenX: sx, screenY: sy, worldX: wx, worldY: wy }};
        }}

        function findNodeAtWorldPos(wx, wy) {{
            for (let i = graphNodes.length - 1; i >= 0; i--) {{
                const n = graphNodes[i];
                if (graphChapterFilter !== 'all' && n.chapterId !== graphChapterFilter) continue;
                const dx = wx - n.x;
                const dy = wy - n.y;
                const hitMargin = n.type === 'card' ? 8 : 6;
                if (dx * dx + dy * dy <= (n.radius + hitMargin) * (n.radius + hitMargin)) {{
                    return n;
                }}
            }}
            return null;
        }}

        function setupGraphInteractions() {{
            const canvas = document.getElementById('graphCanvas');
            const container = document.getElementById('graphCanvasContainer');
            if (!canvas) return;

            // Double click to toggle principle satellite
            canvas.addEventListener('dblclick', e => {{
                const pos = getCanvasMousePos(e);
                const hit = findNodeAtWorldPos(pos.worldX, pos.worldY);
                if (hit && hit.type === 'principle') {{
                    togglePrincipleSatellite(hit.num);
                }}
            }});

            canvas.addEventListener('mousedown', e => {{
                const pos = getCanvasMousePos(e);
                const hit = findNodeAtWorldPos(pos.worldX, pos.worldY);
                mouseMovedDuringClick = false;
                if (hit) {{
                    draggedNode = hit;
                }} else {{
                    isPanning = true;
                    panStartX = pos.screenX - graphCam.x;
                    panStartY = pos.screenY - graphCam.y;
                    container.classList.add('panning');
                }}
            }});

            window.addEventListener('mousemove', e => {{
                if (currentActiveView !== 'graph') return;
                const pos = getCanvasMousePos(e);
                
                if (isPanning) {{
                    mouseMovedDuringClick = true;
                    graphCam.x = pos.screenX - panStartX;
                    graphCam.y = pos.screenY - panStartY;
                }} else if (draggedNode) {{
                    mouseMovedDuringClick = true;
                    draggedNode.x = pos.worldX;
                    draggedNode.y = pos.worldY;
                    draggedNode.vx = 0;
                    draggedNode.vy = 0;
                }} else {{
                    const hit = findNodeAtWorldPos(pos.worldX, pos.worldY);
                    if (hit !== hoveredNode) {{
                        hoveredNode = hit;
                        canvas.style.cursor = hit ? 'pointer' : 'grab';
                    }}
                }}
            }});

            window.addEventListener('mouseup', e => {{
                if (currentActiveView !== 'graph') return;
                if (draggedNode) {{
                    if (!mouseMovedDuringClick) handleNodeClick(draggedNode);
                    draggedNode = null;
                }} else if (isPanning && !mouseMovedDuringClick) {{
                    const pos = getCanvasMousePos(e);
                    const hit = findNodeAtWorldPos(pos.worldX, pos.worldY);
                    if (hit) handleNodeClick(hit);
                }}
                if (isPanning) {{
                    isPanning = false;
                    container.classList.remove('panning');
                }}
            }});

            canvas.addEventListener('wheel', e => {{
                e.preventDefault();
                const pos = getCanvasMousePos(e);
                const zoomFactor = e.deltaY < 0 ? 1.12 : 0.89;
                const newScale = Math.max(0.35, Math.min(2.8, graphCam.scale * zoomFactor));

                graphCam.x -= (pos.worldX * newScale - pos.worldX * graphCam.scale);
                graphCam.y -= (pos.worldY * newScale - pos.worldY * graphCam.scale);
                graphCam.scale = newScale;
            }}, {{ passive: false }});

            // Touch events for mobile
            let touchDist = 0;
            canvas.addEventListener('touchstart', e => {{
                if (e.touches.length === 1) {{
                    const pos = getCanvasMousePos(e);
                    const hit = findNodeAtWorldPos(pos.worldX, pos.worldY);
                    mouseMovedDuringClick = false;
                    if (hit) {{
                        draggedNode = hit;
                    }} else {{
                        isPanning = true;
                        panStartX = pos.screenX - graphCam.x;
                        panStartY = pos.screenY - graphCam.y;
                    }}
                }} else if (e.touches.length === 2) {{
                    isPanning = false;
                    draggedNode = null;
                    const dx = e.touches[0].clientX - e.touches[1].clientX;
                    const dy = e.touches[0].clientY - e.touches[1].clientY;
                    touchDist = Math.sqrt(dx * dx + dy * dy);
                }}
            }}, {{ passive: false }});

            canvas.addEventListener('touchmove', e => {{
                e.preventDefault();
                if (e.touches.length === 1) {{
                    const pos = getCanvasMousePos(e);
                    if (isPanning) {{
                        mouseMovedDuringClick = true;
                        graphCam.x = pos.screenX - panStartX;
                        graphCam.y = pos.screenY - panStartY;
                    }} else if (draggedNode) {{
                        mouseMovedDuringClick = true;
                        draggedNode.x = pos.worldX;
                        draggedNode.y = pos.worldY;
                        draggedNode.vx = 0;
                        draggedNode.vy = 0;
                    }}
                }} else if (e.touches.length === 2) {{
                    const dx = e.touches[0].clientX - e.touches[1].clientX;
                    const dy = e.touches[0].clientY - e.touches[1].clientY;
                    const newDist = Math.sqrt(dx * dx + dy * dy);
                    if (touchDist > 0) {{
                        const factor = newDist / touchDist;
                        graphCam.scale = Math.max(0.35, Math.min(2.8, graphCam.scale * factor));
                    }}
                    touchDist = newDist;
                }}
            }}, {{ passive: false }});

            canvas.addEventListener('touchend', e => {{
                if (draggedNode) {{
                    if (!mouseMovedDuringClick) handleNodeClick(draggedNode);
                    draggedNode = null;
                }} else if (isPanning && !mouseMovedDuringClick && e.changedTouches[0]) {{
                    const pos = getCanvasMousePos(e);
                    const hit = findNodeAtWorldPos(pos.worldX, pos.worldY);
                    if (hit) handleNodeClick(hit);
                }}
                isPanning = false;
                touchDist = 0;
            }});

            window.addEventListener('resize', () => {{
                if (currentActiveView === 'graph') initOrResizeGraph();
            }});
        }}

        function handleNodeClick(node) {{
            if (node.type === 'chapter') {{
                filterGraphChapter(node.chapterId);
            }} else if (node.type === 'principle') {{
                openPrincipleDrawer(node);
            }} else if (node.type === 'card') {{
                openCardDrawer(node);
            }}
        }}

        function setGraphMode(mode) {{
            graphMode = mode;
            const btnP = document.getElementById('btnModePrinciples');
            const btnG = document.getElementById('btnModeGalaxy');
            if (btnP) btnP.classList.toggle('active', mode === 'principles');
            if (btnG) btnG.classList.toggle('active', mode === 'galaxy');
            buildGraphData();
            isGraphSimulating = true;
            updateSatelliteButtonState();
        }}

        function togglePrincipleSatellite(pNum) {{
            if (expandedPrinciples.has(pNum)) {{
                expandedPrinciples.delete(pNum);
            }} else {{
                expandedPrinciples.add(pNum);
            }}
            if (graphMode === 'principles') {{
                buildGraphData();
                isGraphSimulating = true;
            }}
            updateSatelliteButtonState();
        }}

        function toggleCurrentPrincipleSatellite() {{
            if (!selectedNode || selectedNode.type !== 'principle') return;
            if (graphMode === 'galaxy') {{
                setGraphMode('principles');
                expandedPrinciples = new Set([selectedNode.num]);
                buildGraphData();
            }} else {{
                togglePrincipleSatellite(selectedNode.num);
            }}
            updateSatelliteButtonState();
        }}

        function updateSatelliteButtonState() {{
            const satBtn = document.getElementById('btnDrawerSatellite');
            if (!satBtn) return;
            if (!selectedNode || selectedNode.type !== 'principle') {{
                satBtn.style.display = 'none';
                return;
            }}
            satBtn.style.display = 'inline-flex';
            const isExp = graphMode === 'galaxy' || expandedPrinciples.has(selectedNode.num);
            satBtn.innerText = isExp ? '🪐 收合卡牌' : '🪐 展開卡牌';
            satBtn.classList.toggle('active', isExp);
        }}

        function openPrincipleDrawer(node) {{
            selectedNode = node;
            const drawer = document.getElementById('graphDrawer');
            if (!drawer) return;

            document.getElementById('drawerBadge').innerText = `原則 ${{node.num}}`;
            document.getElementById('drawerMeta').innerText = CHAPTER_THEMES[node.chapterId].name;
            document.getElementById('drawerTitle').innerText = node.data.title;
            document.getElementById('drawerMandarin').innerText = node.data.mandarin;
            document.getElementById('drawerEnglish').innerText = node.data.english;
            document.getElementById('drawerTaiwanese').innerText = node.data.taiwanese;

            // Find connected principles
            const connSet = new Set();
            graphEdges.forEach(e => {{
                if (e.source === node.id && e.target.startsWith('p')) {{
                    const targetNum = parseInt(e.target.replace('p', ''));
                    connSet.add(targetNum);
                }} else if (e.target === node.id && e.source.startsWith('p')) {{
                    const sourceNum = parseInt(e.source.replace('p', ''));
                    connSet.add(sourceNum);
                }}
            }});

            const connChipsContainer = document.getElementById('drawerConnectedPrinciples');
            connChipsContainer.innerHTML = '';
            if (connSet.size === 0) {{
                connChipsContainer.innerHTML = `<span style="font-size:0.78rem; color:var(--text-muted);">無跨原則直接重疊</span>`;
            }} else {{
                Array.from(connSet).sort((a, b) => a - b).forEach(num => {{
                    const btn = document.createElement('button');
                    btn.className = 'drawer-conn-chip';
                    btn.innerText = `原則 ${{num}}`;
                    btn.title = `聚焦至原則 ${{num}}`;
                    btn.onclick = () => focusPrincipleInGraph(num);
                    connChipsContainer.appendChild(btn);
                }});
            }}

            // Render list of referenced slide cards
            const refKeys = node.data.ref.split(',').map(r => r.trim()).filter(Boolean);
            document.getElementById('drawerCardsLabel').innerText = `📑 對照簡報卡牌 (${{refKeys.length}} 則，點擊跳轉)`;
            
            const cardsList = document.getElementById('drawerCardsList');
            cardsList.innerHTML = '';
            refKeys.forEach(k => {{
                const s = SLIDES.find(item => item.key === k);
                const mandarinTxt = s ? s.mandarin : '查看簡報筆記項目';
                const item = document.createElement('div');
                item.className = 'drawer-card-item';
                item.title = `點擊跳轉至簡報卡牌 No. ${{k}}`;
                item.innerHTML = `
                    <div class="drawer-card-item-top">
                        <span class="drawer-card-key">No. ${{k}}</span>
                        <span class="drawer-card-jump">跳轉卡牌 ↗</span>
                    </div>
                    <div class="drawer-card-mandarin">${{mandarinTxt}}</div>
                `;
                item.onclick = () => jumpToSlideCard(k);
                cardsList.appendChild(item);
            }});

            updateSatelliteButtonState();
            drawer.classList.add('open');
        }}

        function openCardDrawer(node) {{
            selectedNode = node;
            const drawer = document.getElementById('graphDrawer');
            if (!drawer) return;

            document.getElementById('drawerBadge').innerText = `卡牌 No. ${{node.key}}`;
            const chapName = CHAPTER_THEMES[node.chapterId] ? CHAPTER_THEMES[node.chapterId].name : '簡報卡牌';
            document.getElementById('drawerMeta').innerText = `Slide ${{node.data.slide}} · ${{chapName}}`;
            document.getElementById('drawerTitle').innerText = node.isBridge ? `跨界交匯卡牌 No. ${{node.key}} (關聯 ${{node.principles.length}} 原則)` : `簡報卡牌 No. ${{node.key}}`;

            document.getElementById('drawerMandarin').innerText = node.data.mandarin || '無中文潤稿';
            document.getElementById('drawerEnglish').innerText = node.data.english || node.data.original || '';
            document.getElementById('drawerTaiwanese').innerText = node.data.taiwanese || '無台文潤稿';

            // Connected principles
            const connChipsContainer = document.getElementById('drawerConnectedPrinciples');
            connChipsContainer.innerHTML = '';
            node.principles.forEach(pNum => {{
                const pObj = PRINCIPLES.find(p => p.num === pNum);
                const btn = document.createElement('button');
                btn.className = 'drawer-conn-chip';
                btn.innerText = `原則 ${{pNum}}${{pObj ? ' ' + pObj.title.slice(0, 7) + '..' : ''}}`;
                btn.title = `聚焦至原則 ${{pNum}}`;
                btn.onclick = () => focusPrincipleInGraph(pNum);
                connChipsContainer.appendChild(btn);
            }});

            document.getElementById('drawerCardsLabel').innerText = `🔗 原文摘錄與導航`;
            const cardsList = document.getElementById('drawerCardsList');
            cardsList.innerHTML = `
                <div class="drawer-card-item" onclick="jumpToSlideCard('${{node.key}}')" style="background: rgba(99, 102, 241, 0.1); border-color: var(--primary);">
                    <div class="drawer-card-item-top">
                        <span class="drawer-card-key">🗂️ 前往「簡報卡牌」視角檢視完整卡牌</span>
                        <span class="drawer-card-jump">點擊跳轉 ↗</span>
                    </div>
                    <div class="drawer-card-mandarin" style="color: var(--text-muted); font-size: 0.78rem;">
                        <b>投影片原文：</b>${{node.data.original || ''}}
                    </div>
                </div>
            `;

            updateSatelliteButtonState();
            drawer.classList.add('open');
        }}

        function closeGraphDrawer() {{
            const drawer = document.getElementById('graphDrawer');
            if (drawer) drawer.classList.remove('open');
            selectedNode = null;
            updateSatelliteButtonState();
        }}

        function focusPrincipleInGraph(num) {{
            const target = graphNodes.find(n => n.id === `p${{num}}`);
            if (target) {{
                graphCam.x = -target.x * graphCam.scale;
                graphCam.y = -target.y * graphCam.scale;
                openPrincipleDrawer(target);
            }}
        }}

        function filterGraphChapter(chapterId) {{
            graphChapterFilter = chapterId;
            const filterButtons = document.querySelectorAll('#graphFilterChips .filter-tag');
            filterButtons.forEach(btn => btn.classList.remove('active'));
            if (event && event.target && event.target.classList.contains('filter-tag')) {{
                event.target.classList.add('active');
            }} else {{
                const idx = chapterId === 'all' ? 0 : parseInt(chapterId);
                if (filterButtons[idx]) filterButtons[idx].classList.add('active');
            }}

            if (chapterId !== 'all') {{
                const chapNode = graphNodes.find(n => n.id === `c${{chapterId}}`);
                if (chapNode) {{
                    graphCam.x = -chapNode.x * graphCam.scale;
                    graphCam.y = -chapNode.y * graphCam.scale;
                }}
            }}
        }}

        function resetGraphView() {{
            graphCam = {{ x: 0, y: 0, scale: 1.0 }};
            selectedNode = null;
            closeGraphDrawer();
        }}

        function toggleGraphSimulation() {{
            isGraphSimulating = !isGraphSimulating;
            const btn = document.getElementById('btnGraphFreeze');
            if (btn) {{
                btn.innerText = isGraphSimulating ? '⏸️ 暫停模擬' : '▶️ 繼續模擬';
            }}
        }}

        function stepPhysics() {{
            if (!isGraphSimulating) return;

            const kRepel = 2200;
            const maxRepelDist = 180;
            const maxRepelDistSq = maxRepelDist * maxRepelDist;

            // 1. Repulsion with spatial distance pruning
            for (let i = 0; i < graphNodes.length; i++) {{
                const n1 = graphNodes[i];
                for (let j = i + 1; j < graphNodes.length; j++) {{
                    const n2 = graphNodes[j];
                    const dx = n2.x - n1.x;
                    const dy = n2.y - n1.y;
                    const distSq = dx * dx + dy * dy;

                    // Prune far card-to-card pairs for 60fps performance
                    if (distSq > maxRepelDistSq && n1.type === 'card' && n2.type === 'card') continue;
                    if (distSq < 1) continue;

                    const dist = Math.sqrt(distSq);
                    const force = (kRepel * (n1.mass * n2.mass)) / (distSq + 240);
                    const fx = (dx / dist) * force;
                    const fy = (dy / dist) * force;

                    n1.vx -= fx / n1.mass;
                    n1.vy -= fy / n1.mass;
                    n2.vx += fx / n2.mass;
                    n2.vy += fy / n2.mass;
                }}
            }}

            // 2. Spring force along edges
            graphEdges.forEach(e => {{
                const s = graphNodes.find(n => n.id === e.source);
                const t = graphNodes.find(n => n.id === e.target);
                if (!s || !t) return;
                const dx = t.x - s.x;
                const dy = t.y - s.y;
                const dist = Math.sqrt(dx * dx + dy * dy) || 1;
                const delta = dist - e.distance;
                const force = delta * e.strength;
                const fx = (dx / dist) * force;
                const fy = (dy / dist) * force;

                s.vx += fx / s.mass;
                s.vy += fy / s.mass;
                t.vx -= fx / t.mass;
                t.vy -= fy / t.mass;
            }});

            // 3. Center gravity and damping
            graphNodes.forEach(node => {{
                if (node === draggedNode) return;
                
                const pull = node.type === 'card' ? 0.0008 : 0.0025;
                node.vx += (-node.x) * pull;
                node.vy += (-node.y) * pull;

                node.vx *= 0.86;
                node.vy *= 0.86;

                const speed = Math.sqrt(node.vx * node.vx + node.vy * node.vy);
                if (speed > 8) {{
                    node.vx = (node.vx / speed) * 8;
                    node.vy = (node.vy / speed) * 8;
                }}

                node.x += node.vx;
                node.y += node.vy;
            }});
        }}

        function renderGraph() {{
            const canvas = document.getElementById('graphCanvas');
            if (!canvas) return;
            const ctx = canvas.getContext('2d');
            const dpr = window.devicePixelRatio || 1;
            const w = canvas.width / dpr;
            const h = canvas.height / dpr;

            ctx.clearRect(0, 0, canvas.width, canvas.height);

            ctx.save();
            ctx.scale(dpr, dpr);
            ctx.translate(w / 2 + graphCam.x, h / 2 + graphCam.y);
            ctx.scale(graphCam.scale, graphCam.scale);

            const activeNode = selectedNode || hoveredNode;
            const connectedNodeIds = new Set();
            if (activeNode) {{
                connectedNodeIds.add(activeNode.id);
                graphEdges.forEach(e => {{
                    if (e.source === activeNode.id) connectedNodeIds.add(e.target);
                    if (e.target === activeNode.id) connectedNodeIds.add(e.source);
                }});
            }}

            const isLight = document.body.classList.contains('light-mode');

            // Draw Edges
            graphEdges.forEach(e => {{
                const s = graphNodes.find(n => n.id === e.source);
                const t = graphNodes.find(n => n.id === e.target);
                if (!s || !t) return;

                const isConnected = activeNode && (e.source === activeNode.id || e.target === activeNode.id);
                let alpha = 0.28;
                if (activeNode) {{
                    alpha = isConnected ? 0.95 : 0.06;
                }} else if (graphChapterFilter !== 'all') {{
                    const inChap = (s.chapterId === graphChapterFilter || t.chapterId === graphChapterFilter);
                    alpha = inChap ? 0.5 : 0.06;
                }}

                ctx.save();
                ctx.globalAlpha = alpha;

                if (e.type === 'chapter') {{
                    ctx.strokeStyle = e.color;
                    ctx.lineWidth = isConnected ? 2.8 : 1.2;
                    ctx.setLineDash([4, 4]);
                }} else if (e.type === 'shared') {{
                    ctx.strokeStyle = isLight ? '#64748b' : '#94a3b8';
                    ctx.lineWidth = isConnected ? 3.2 : (1.4 + Math.min(2.5, e.sharedCount * 0.7));
                    ctx.setLineDash([]);
                }} else if (e.type === 'synergy') {{
                    ctx.strokeStyle = '#c084fc';
                    ctx.lineWidth = isConnected ? 2.5 : 1.2;
                    ctx.setLineDash([6, 3]);
                }} else if (e.type === 'card') {{
                    ctx.strokeStyle = isConnected ? e.color : (isLight ? '#cbd5e1' : '#475569');
                    ctx.lineWidth = isConnected ? 2.0 : 0.8;
                    ctx.setLineDash([]);
                }}

                ctx.beginPath();
                ctx.moveTo(s.x, s.y);
                ctx.lineTo(t.x, t.y);
                ctx.stroke();

                // Draw badge on shared edges if sharedCount >= 2 and active
                if (e.type === 'shared' && e.sharedCount >= 2 && (isConnected || !activeNode)) {{
                    const mx = (s.x + t.x) / 2;
                    const my = (s.y + t.y) / 2;
                    ctx.fillStyle = isLight ? 'rgba(255,255,255,0.9)' : 'rgba(15,23,42,0.9)';
                    ctx.beginPath();
                    ctx.arc(mx, my, 8, 0, Math.PI * 2);
                    ctx.fill();
                    ctx.strokeStyle = '#94a3b8';
                    ctx.stroke();
                    ctx.fillStyle = isLight ? '#0f172a' : '#f8fafc';
                    ctx.font = 'bold 8px Outfit, sans-serif';
                    ctx.textAlign = 'center';
                    ctx.textBaseline = 'middle';
                    ctx.fillText(e.sharedCount, mx, my);
                }}

                ctx.restore();
            }});

            // Draw Nodes
            graphNodes.forEach(node => {{
                const isFiltered = (graphChapterFilter !== 'all' && node.chapterId !== graphChapterFilter);
                let alpha = 1.0;
                if (isFiltered) {{
                    alpha = 0.15;
                }} else if (activeNode) {{
                    alpha = connectedNodeIds.has(node.id) ? 1.0 : 0.18;
                }}

                const isSelected = selectedNode && selectedNode.id === node.id;
                const isHovered = hoveredNode && hoveredNode.id === node.id;

                ctx.save();
                ctx.globalAlpha = alpha;

                if (node.type === 'chapter') {{
                    // Chapter Hub Node
                    ctx.beginPath();
                    ctx.arc(node.x, node.y, node.radius + (isHovered ? 4 : 0), 0, Math.PI * 2);
                    ctx.fillStyle = node.color;
                    ctx.fill();

                    // Outer ring
                    ctx.lineWidth = 3;
                    ctx.strokeStyle = isLight ? '#ffffff' : 'rgba(255,255,255,0.85)';
                    ctx.stroke();

                    // Label
                    ctx.fillStyle = '#ffffff';
                    ctx.font = 'bold 12px Outfit, sans-serif';
                    ctx.textAlign = 'center';
                    ctx.textBaseline = 'middle';
                    ctx.fillText(`Ch.0${{node.num}}`, node.x, node.y - 4);

                    ctx.font = '9px "Noto Sans TC", sans-serif';
                    ctx.fillText(node.title.slice(0, 4), node.x, node.y + 9);
                }} else if (node.type === 'principle') {{
                    // Principle Node
                    const isExp = graphMode === 'galaxy' || expandedPrinciples.has(node.num);

                    if (isSelected || isHovered) {{
                        ctx.beginPath();
                        ctx.arc(node.x, node.y, node.radius + 6, 0, Math.PI * 2);
                        ctx.fillStyle = isLight ? 'rgba(99, 102, 241, 0.2)' : 'rgba(99, 102, 241, 0.35)';
                        ctx.fill();
                    }}

                    // Satellite orbit indicator if expanded
                    if (isExp) {{
                        ctx.beginPath();
                        ctx.arc(node.x, node.y, node.radius + 5, 0, Math.PI * 2);
                        ctx.strokeStyle = node.color;
                        ctx.lineWidth = 1.2;
                        ctx.setLineDash([3, 3]);
                        ctx.stroke();
                        ctx.setLineDash([]);
                    }}

                    ctx.beginPath();
                    ctx.arc(node.x, node.y, node.radius, 0, Math.PI * 2);
                    ctx.fillStyle = isLight ? '#ffffff' : 'rgba(30, 41, 59, 0.95)';
                    ctx.fill();

                    ctx.lineWidth = isSelected ? 3.5 : (isHovered ? 2.5 : 1.8);
                    ctx.strokeStyle = node.color;
                    ctx.stroke();

                    // Principle number
                    ctx.fillStyle = isLight ? node.color : '#f8fafc';
                    ctx.font = 'bold 12px Outfit, sans-serif';
                    ctx.textAlign = 'center';
                    ctx.textBaseline = 'middle';
                    ctx.fillText(`P${{node.num}}`, node.x, node.y);

                    // Short title under node
                    if (alpha > 0.4) {{
                        ctx.fillStyle = isLight ? '#334155' : '#cbd5e1';
                        ctx.font = '10px "Noto Sans TC", sans-serif';
                        ctx.textBaseline = 'top';
                        const displayTitle = node.title.length > 7 ? node.title.slice(0, 6) + '..' : node.title;
                        ctx.fillText(displayTitle, node.x, node.y + node.radius + 4);
                    }}
                }} else if (node.type === 'card') {{
                    // Card Satellite Node
                    if (isSelected || isHovered) {{
                        ctx.beginPath();
                        ctx.arc(node.x, node.y, node.radius + 4, 0, Math.PI * 2);
                        ctx.fillStyle = node.isBridge ? 'rgba(251, 191, 36, 0.35)' : 'rgba(99, 102, 241, 0.3)';
                        ctx.fill();
                    }}

                    ctx.beginPath();
                    ctx.arc(node.x, node.y, node.radius + (isHovered ? 2 : 0), 0, Math.PI * 2);
                    ctx.fillStyle = node.isBridge ? '#fbbf24' : node.color;
                    ctx.fill();

                    ctx.lineWidth = isSelected ? 2.5 : (node.isBridge ? 2 : 1.2);
                    ctx.strokeStyle = isLight ? '#ffffff' : '#0f172a';
                    ctx.stroke();

                    // Inner dot for bridge card
                    if (node.isBridge) {{
                        ctx.beginPath();
                        ctx.arc(node.x, node.y, 2.2, 0, Math.PI * 2);
                        ctx.fillStyle = '#ffffff';
                        ctx.fill();
                    }}

                    // Card key text when zoomed in or hovered
                    if (graphCam.scale >= 0.85 || isHovered || isSelected) {{
                        ctx.fillStyle = isLight ? '#334155' : '#e2e8f0';
                        ctx.font = 'bold 8px Outfit, sans-serif';
                        ctx.textAlign = 'center';
                        ctx.textBaseline = 'top';
                        ctx.fillText(node.key, node.x, node.y + node.radius + 3);
                    }}
                }}

                ctx.restore();
            }});

            // Draw hover tooltip for card node
            if (hoveredNode && hoveredNode.type === 'card') {{
                const tipText = hoveredNode.data.mandarin ? hoveredNode.data.mandarin.slice(0, 24) + '..' : `卡牌 No. ${{hoveredNode.key}}`;
                const tipTitle = `No. ${{hoveredNode.key}}${{hoveredNode.isBridge ? ' (跨界交匯卡)' : ''}}`;
                
                ctx.save();
                ctx.font = 'bold 10px Outfit, sans-serif';
                const titleWidth = ctx.measureText(tipTitle).width;
                ctx.font = '9px "Noto Sans TC", sans-serif';
                const textWidth = ctx.measureText(tipText).width;
                const boxW = Math.max(titleWidth, textWidth) + 20;
                const boxH = 36;
                const boxX = hoveredNode.x - boxW / 2;
                const boxY = hoveredNode.y - hoveredNode.radius - boxH - 8;

                ctx.fillStyle = isLight ? 'rgba(255, 255, 255, 0.96)' : 'rgba(15, 23, 42, 0.95)';
                ctx.shadowColor = 'rgba(0,0,0,0.3)';
                ctx.shadowBlur = 8;
                ctx.beginPath();
                ctx.roundRect(boxX, boxY, boxW, boxH, 6);
                ctx.fill();
                ctx.shadowBlur = 0;
                ctx.strokeStyle = hoveredNode.isBridge ? '#fbbf24' : hoveredNode.color;
                ctx.lineWidth = 1.2;
                ctx.stroke();

                ctx.fillStyle = hoveredNode.isBridge ? (isLight ? '#b45309' : '#fbbf24') : hoveredNode.color;
                ctx.font = 'bold 9.5px Outfit, sans-serif';
                ctx.textAlign = 'left';
                ctx.textBaseline = 'top';
                ctx.fillText(tipTitle, boxX + 9, boxY + 5);

                ctx.fillStyle = isLight ? '#334155' : '#cbd5e1';
                ctx.font = '8.5px "Noto Sans TC", sans-serif';
                ctx.fillText(tipText, boxX + 9, boxY + 19);
                ctx.restore();
            }}

            ctx.restore();
        }}

        function graphAnimationLoop() {{
            if (currentActiveView === 'graph') {{
                stepPhysics();
                renderGraph();
                graphAnimationId = requestAnimationFrame(graphAnimationLoop);
            }} else {{
                graphAnimationId = null;
            }}
        }}
    </script>
</body>
</html>
"""
    
    with open("index.html", "w", encoding="utf-8") as f:
        f.write(html_content)
        
    print(f"Saved regenerated HTML card web app: index.html")
    
    # Generate ODT handbook as well
    odt_path = "工作原則整理與潤稿.odt"
    generate_odt_handbook(principles, final_items, odt_path, chapters=chapters)
    
    print("\nWorkflow completed successfully!")

if __name__ == "__main__":
    main()
